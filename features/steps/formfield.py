"""Step implementations for legacy form-field features."""

from behave import given, then, when
from behave.runner import Context

from docx import Document
from docx.enum.text import WD_FORM_FIELD_TYPE, WD_TEXT_FORM_FIELD_TYPE

from helpers import saved_docx_path, test_docx

# given ===================================================


@given("a document containing legacy form fields")
def given_a_document_containing_legacy_form_fields(context: Context):
    context.document = Document(test_docx("frm-form-fields"))


# when ====================================================


@when("I assign a new value to each form field")
def when_I_assign_a_new_value_to_each_form_field(context: Context):
    fields = {f.name: f for f in context.document.form_fields}
    fields["Applicant"].value = "Carol Chen"
    fields["Agreed"].value = False
    fields["Tier"].value = "Lifetime"
    fields["Referrer"].value = "Dave Doyle"


# then ====================================================


@then("document.form_fields has an item for each field in the document")
def then_document_form_fields_has_an_item_for_each_field(context: Context):
    names = [f.name for f in context.document.form_fields]
    assert names == ["Applicant", "Agreed", "Tier", "Referrer"], f"got {names}"


@then("each form field knows its name and kind")
def then_each_form_field_knows_its_name_and_kind(context: Context):
    kinds = {f.name: f.type for f in context.document.form_fields}
    assert kinds == {
        "Applicant": WD_FORM_FIELD_TYPE.TEXT,
        "Agreed": WD_FORM_FIELD_TYPE.CHECK_BOX,
        "Tier": WD_FORM_FIELD_TYPE.DROP_DOWN,
        "Referrer": WD_FORM_FIELD_TYPE.TEXT,
    }, f"got {kinds}"


@then("paragraph.form_fields has the fields of that paragraph")
def then_paragraph_form_fields_has_the_fields_of_that_paragraph(context: Context):
    paragraph = context.document.paragraphs[1]
    names = [f.name for f in paragraph.form_fields]
    assert names == ["Applicant"], f"got {names}"
    # -- the field in the table cell belongs to that cell's paragraph, not the body --
    cell_paragraph = context.document.tables[0].cell(0, 1).paragraphs[0]
    cell_names = [f.name for f in cell_paragraph.form_fields]
    assert cell_names == ["Referrer"], f"got {cell_names}"


@then("a text form field knows the text it holds")
def then_a_text_form_field_knows_the_text_it_holds(context: Context):
    field = _field(context, "Applicant")
    assert field.value == "Alice Ashby", f"got {field.value!r}"


@then("a check box form field knows whether it is ticked")
def then_a_check_box_form_field_knows_whether_it_is_ticked(context: Context):
    field = _field(context, "Agreed")
    assert field.value is True, f"got {field.value!r}"


@then("a drop-down form field knows the entry selected and the entries offered")
def then_a_drop_down_form_field_knows_its_entries(context: Context):
    field = _field(context, "Tier")
    assert field.value == "Premium", f"got {field.value!r}"
    assert field.items == ("Standard", "Premium", "Lifetime"), f"got {field.items}"


@then("a form field knows its help text, status text and enabled state")
def then_a_form_field_knows_its_help_text_status_text_and_enabled_state(context: Context):
    field = _field(context, "Applicant")
    assert field.help_text == "Your full name", f"got {field.help_text!r}"
    assert field.status_text == "Applicant name", f"got {field.status_text!r}"
    assert field.enabled is True, f"got {field.enabled!r}"


@then("a text form field knows its default, maximum length and text type")
def then_a_text_form_field_knows_its_default_max_length_and_text_type(context: Context):
    field = _field(context, "Applicant")
    assert field.default == "Type your name", f"got {field.default!r}"
    assert field.max_length == 30, f"got {field.max_length!r}"
    assert field.text_type == WD_TEXT_FORM_FIELD_TYPE.REGULAR_TEXT, f"got {field.text_type!r}"


@then("the reloaded document reports the new values")
def then_the_reloaded_document_reports_the_new_values(context: Context):
    document = Document(saved_docx_path)
    values = {f.name: f.value for f in document.form_fields}
    assert values == {
        "Applicant": "Carol Chen",
        "Agreed": False,
        "Tier": "Lifetime",
        "Referrer": "Dave Doyle",
    }, f"got {values}"


def _field(context: Context, name: str):
    return next(f for f in context.document.form_fields if f.name == name)
