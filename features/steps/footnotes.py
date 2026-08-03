"""Step implementations for document footnote features."""

from behave import given, then, when
from behave.runner import Context

from docx import Document
from docx.oxml.ns import qn

from helpers import saved_docx_path, test_docx

# given ===================================================


@given("a document having no footnotes part")
def given_a_document_having_no_footnotes_part(context: Context):
    context.document = Document(test_docx("doc-default"))


# when ====================================================


@when("I add a footnote and reference it from a run")
def when_I_add_a_footnote_and_reference_it_from_a_run(context: Context):
    document = context.document
    paragraph = document.add_paragraph("The claim is well established")
    footnote = document.footnotes.add_footnote("See Smith (2019), p. 42.")
    context.run = paragraph.add_run()
    context.run.add_footnote_reference(footnote)


# then ====================================================


@then("document.footnotes has zero footnotes")
def then_document_footnotes_has_zero_footnotes(context: Context):
    count = len(context.document.footnotes)
    assert count == 0, f"expected 0 footnotes, got {count}"


@then("document.footnotes has one footnote")
def then_document_footnotes_has_one_footnote(context: Context):
    count = len(context.document.footnotes)
    assert count == 1, f"expected 1 footnote, got {count}"


@then("the footnote knows its id and text")
def then_the_footnote_knows_its_id_and_text(context: Context):
    footnote = list(context.document.footnotes)[0]
    # -- ids -1 and 0 are reserved for the separators, so an author's start at 1 --
    assert footnote.footnote_id == 1, f"got {footnote.footnote_id}"
    assert footnote.text == "See Smith (2019), p. 42.", f"got {footnote.text!r}"


@then("the run holds a reference to the footnote")
def then_the_run_holds_a_reference_to_the_footnote(context: Context):
    ref = context.run._r.find(qn("w:footnoteReference"))
    assert ref is not None, "run holds no `w:footnoteReference`"
    assert ref.get(qn("w:id")) == "1", f"got {ref.get(qn('w:id'))!r}"


@then("the reloaded document reports the footnote")
def then_the_reloaded_document_reports_the_footnote(context: Context):
    document = Document(saved_docx_path)
    texts = [f.text for f in document.footnotes]
    assert texts == ["See Smith (2019), p. 42."], f"got {texts}"


@then("the reloaded document keeps the separator footnotes out of the collection")
def then_the_reloaded_document_keeps_the_separators_out(context: Context):
    footnotes = Document(saved_docx_path).footnotes
    assert footnotes.get(-1) is None, "separator footnote is in the collection"
    assert footnotes.get(0) is None, "continuation separator is in the collection"
    ids = [f.footnote_id for f in footnotes]
    assert ids == [1], f"got {ids}"
