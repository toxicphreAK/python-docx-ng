"""Step implementations for structured document tag (content control) features."""

from behave import given, then
from behave.runner import Context

from docx import Document
from docx.enum.text import WD_CONTENT_CONTROL_TYPE

from helpers import test_docx

# given ===================================================


@given("a document containing content controls")
def given_a_document_containing_content_controls(context: Context):
    context.document = Document(test_docx("sdt-content-controls"))


# then ====================================================


@then("document.paragraphs includes the paragraphs inside the content controls")
def then_document_paragraphs_includes_content_control_paragraphs(context: Context):
    texts = [p.text for p in context.document.paragraphs]
    assert texts == [
        "Contract summary",
        "Acme Corporation",
        "Click here to enter text.",
        "Signed on 3 August 2026 by both parties.",
        "☒ Terms accepted",
        "End of summary.",
    ], f"got {texts}"


@then("document.tables includes the table inside a content control")
def then_document_tables_includes_the_wrapped_table(context: Context):
    tables = context.document.tables
    assert len(tables) == 1, f"expected 1 table, got {len(tables)}"
    assert tables[0].cell(1, 0).text == "Consulting"


@then("paragraph.runs includes the runs inside a run-level content control")
def then_paragraph_runs_includes_wrapped_runs(context: Context):
    paragraph = context.document.paragraphs[3]
    texts = [r.text for r in paragraph.runs]
    assert texts == ["Signed on ", "3 August 2026", " by both parties."], f"got {texts}"


@then("paragraph.text includes the text of a run-level content control")
def then_paragraph_text_includes_wrapped_text(context: Context):
    paragraph = context.document.paragraphs[3]
    assert paragraph.text == "Signed on 3 August 2026 by both parties."


@then("document.content_controls has an item for each control in the document")
def then_document_content_controls_has_an_item_for_each(context: Context):
    tags = [cc.tag for cc in context.document.content_controls]
    assert tags == [
        "client_name",
        "signature",
        "signed_on",
        "agreed",
        "line_items",
    ], f"got {tags}"


@then("each content control knows its tag, alias and type")
def then_each_content_control_knows_its_identity(context: Context):
    client_name, signature, signed_on, agreed, _ = context.document.content_controls

    assert client_name.alias == "Client name"
    assert client_name.id == 101
    assert client_name.type == WD_CONTENT_CONTROL_TYPE.TEXT
    assert signature.type == WD_CONTENT_CONTROL_TYPE.RICH_TEXT
    assert signed_on.type == WD_CONTENT_CONTROL_TYPE.DATE
    assert agreed.type == WD_CONTENT_CONTROL_TYPE.CHECKBOX


@then("a content control knows when it is showing its placeholder")
def then_a_content_control_knows_it_is_showing_its_placeholder(context: Context):
    client_name, signature = context.document.content_controls[:2]

    assert signature.showing_placeholder is True
    assert client_name.showing_placeholder is False


@then("a content control knows the text it contains")
def then_a_content_control_knows_its_text(context: Context):
    client_name, _, signed_on, _, line_items = context.document.content_controls

    assert client_name.text == "Acme Corporation"
    assert signed_on.text == "3 August 2026"
    assert line_items.text == "Item\nAmount\nConsulting\n1000"
