"""Unit test suite for the docx.fields module."""

from __future__ import annotations

import io
from typing import cast

import pytest

import docx
from docx import fields
from docx import types as t
from docx.fields import Field
from docx.oxml.ns import nsdecls
from docx.oxml.parser import parse_xml
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph


def _p(inner_xml: str) -> CT_P:
    """A `w:p` element containing `inner_xml`.

    Field markup is written out rather than expressed in CXEL because a field
    instruction carries the leading and trailing spaces Word writes, and the CXEL
    attribute grammar cannot express them.
    """
    return cast(CT_P, parse_xml(f"<w:p {nsdecls('w')}>{inner_xml}</w:p>"))


def _run(inner_xml: str) -> str:
    return f"<w:r>{inner_xml}</w:r>"


def _fldChar(fldCharType: str, dirty: bool = False) -> str:
    attrs = f'w:fldCharType="{fldCharType}"' + (' w:dirty="true"' if dirty else "")
    return _run(f"<w:fldChar {attrs}/>")


def _instr(text: str) -> str:
    return _run(f'<w:instrText xml:space="preserve">{text}</w:instrText>')


def _text(text: str) -> str:
    return _run(f"<w:t>{text}</w:t>")


# -- the complex field Word writes for a page number --
_PAGE_FIELD = (
    _fldChar("begin") + _instr(" PAGE ") + _fldChar("separate") + _text("7") + _fldChar("end")
)


class DescribeFieldReading:
    """Unit-test suite for reading fields out of a document."""

    def it_reads_a_complex_field(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(_p(_PAGE_FIELD), fake_parent)

        (field,) = paragraph.fields

        assert field.type == "PAGE"
        assert field.instruction == " PAGE "
        assert field.result_text == "7"
        assert field.is_simple is False

    def it_reads_a_simple_field(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(
            _p('<w:fldSimple w:instr=" PAGE "><w:r><w:t>7</w:t></w:r></w:fldSimple>'),
            fake_parent,
        )

        (field,) = paragraph.fields

        assert field.type == "PAGE"
        assert field.instruction == " PAGE "
        assert field.result_text == "7"
        assert field.is_simple is True

    def it_joins_an_instruction_word_split_across_runs(self, fake_parent: t.ProvidesStoryPart):
        """Word routinely splits an instruction; a per-run reader sees only fragments."""
        paragraph = Paragraph(
            _p(
                _fldChar("begin")
                + _instr(" TOC ")
                + _instr("\\o ")
                + _instr("&quot;1-3&quot; ")
                + _fldChar("end")
            ),
            fake_parent,
        )

        (field,) = paragraph.fields

        assert field.type == "TOC"
        assert field.instruction == ' TOC \\o "1-3" '

    def it_reads_a_field_nested_in_another_fields_result(self, fake_parent: t.ProvidesStoryPart):
        """A TOC result is full of PAGEREF fields; each is a field in its own right."""
        paragraph = Paragraph(
            _p(
                _fldChar("begin")
                + _instr(" TOC ")
                + _fldChar("separate")
                + _text("Intro")
                + _fldChar("begin")
                + _instr(" PAGEREF _Toc1 ")
                + _fldChar("separate")
                + _text("3")
                + _fldChar("end")
                + _fldChar("end")
            ),
            fake_parent,
        )

        toc, pageref = paragraph.fields

        assert (toc.type, pageref.type) == ("TOC", "PAGEREF")
        assert pageref.result_text == "3"
        # -- the nested field's result is part of what the TOC entry displays --
        assert toc.result_text == "Intro3"
        # -- but the nested field's instruction is not --
        assert toc.instruction == " TOC "

    def it_knows_whether_a_field_is_marked_dirty(self, fake_parent: t.ProvidesStoryPart):
        clean = Paragraph(_p(_PAGE_FIELD), fake_parent).fields[0]
        dirty = Paragraph(
            _p(_fldChar("begin", dirty=True) + _instr(" PAGE ") + _fldChar("end")),
            fake_parent,
        ).fields[0]

        assert clean.dirty is False
        assert dirty.dirty is True

    def it_skips_a_field_with_no_end_marker(self, fake_parent: t.ProvidesStoryPart):
        """Such documents exist; reading the well-formed fields beats refusing them."""
        paragraph = Paragraph(
            _p(_fldChar("begin") + _instr(" PAGE ") + _text("orphan")), fake_parent
        )

        assert paragraph.fields == []

    def it_reads_a_field_spanning_several_paragraphs_at_document_level(self):
        """A table of contents begins in one paragraph and ends several later."""
        document = docx.Document()
        body = document.element.body
        body.insert(
            0,
            parse_xml(
                f"<w:sdt {nsdecls('w')}><w:sdtContent>"
                f"<w:p>{_fldChar('begin')}{_instr(' TOC ')}"
                f"{_fldChar('separate')}{_text('Intro')}</w:p>"
                f"<w:p>{_text('More')}{_fldChar('end')}</w:p>"
                "</w:sdtContent></w:sdt>"
            ),
        )

        (field,) = document.fields
        assert field.type == "TOC"
        assert field.result_text == "IntroMore"
        # -- one paragraph cannot see the end marker, so it reports no field --
        assert document.paragraphs[0].fields == []

    def it_reads_every_field_in_the_document_body(self):
        document = docx.Document()
        document.add_paragraph().add_field(fields.page_number())
        document.add_paragraph().add_field(fields.page_count())

        assert [f.type for f in document.fields] == ["PAGE", "NUMPAGES"]

    @pytest.mark.parametrize(
        ("instruction", "expected_type"),
        [
            (" PAGE ", "PAGE"),
            (" toc ", "TOC"),
            (" REF _Ref1 \\h ", "REF"),
            ("", None),
            (" \\* MERGEFORMAT ", None),
        ],
    )
    def it_takes_the_field_type_from_the_first_token(
        self, instruction: str, expected_type: str | None, fake_parent: t.ProvidesStoryPart
    ):
        fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="x"/>')
        field = Field(cast("t.Any", fldSimple), fake_parent, instruction)

        assert field.type == expected_type

    def it_includes_a_simple_fields_result_in_the_paragraph_text(
        self, fake_parent: t.ProvidesStoryPart
    ):
        """A page number is displayed text; skipping it drops it from `.text`."""
        paragraph = Paragraph(
            _p(
                _text("Page ")
                + '<w:fldSimple w:instr=" PAGE "><w:r><w:t>7</w:t></w:r></w:fldSimple>'
            ),
            fake_parent,
        )

        assert paragraph.text == "Page 7"

    def it_never_reports_a_field_instruction_as_document_text(
        self, fake_parent: t.ProvidesStoryPart
    ):
        paragraph = Paragraph(_p(_PAGE_FIELD), fake_parent)

        assert paragraph.text == "7", "the instruction is not displayed text"


class DescribeAddField:
    """Unit-test suite for `docx.text.paragraph.Paragraph.add_field`."""

    def it_adds_a_complex_field_by_default(self):
        paragraph = docx.Document().add_paragraph()

        field = paragraph.add_field("PAGE")

        assert field.type == "PAGE"
        assert field.is_simple is False
        assert field.dirty is True
        assert [c.tag.split("}")[1] for c in paragraph._p] == ["r", "r", "r"]
        assert len(paragraph._p.xpath(".//w:fldChar")) == 2

    def it_writes_the_instruction_with_space_preserved(self):
        """Without `xml:space` the instruction's delimiting spaces can be collapsed."""
        paragraph = docx.Document().add_paragraph()

        paragraph.add_field("PAGE")

        (instrText,) = paragraph._p.xpath(".//w:instrText")
        assert instrText.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"

    def it_adds_a_simple_field_when_asked(self):
        paragraph = docx.Document().add_paragraph()

        field = paragraph.add_field("PAGE", simple=True, result="7")

        assert field.is_simple is True
        assert field.result_text == "7"
        assert paragraph.text == "7"

    def it_normalizes_the_spacing_around_the_instruction(self):
        paragraph = docx.Document().add_paragraph()

        field = paragraph.add_field("  PAGE  ")

        assert field.instruction == " PAGE "
        assert paragraph.fields[0].instruction == " PAGE "

    def it_can_leave_the_field_clean(self):
        paragraph = docx.Document().add_paragraph()

        field = paragraph.add_field("PAGE", dirty=False)

        assert field.dirty is False

    def it_rejects_a_cached_result_for_a_complex_field(self):
        paragraph = docx.Document().add_paragraph()

        with pytest.raises(ValueError, match="applies only to a simple field"):
            paragraph.add_field("PAGE", result="7")

    def it_can_toggle_the_dirty_flag_afterwards(self):
        paragraph = docx.Document().add_paragraph()
        field = paragraph.add_field("PAGE", dirty=False)

        field.dirty = True

        assert paragraph.fields[0].dirty is True

    def it_adds_the_field_after_the_existing_content(self):
        paragraph = docx.Document().add_paragraph("Page ")

        paragraph.add_field(fields.page_number())
        paragraph.add_run(" of end")

        assert [c.tag.split("}")[1] for c in paragraph._p] == ["r", "r", "r", "r", "r"]
        assert paragraph.text == "Page  of end", "the field has no result yet"

    def it_round_trips_a_field_through_a_saved_document(self):
        document = docx.Document()
        document.add_paragraph().add_field(fields.table_of_contents())
        document.settings.update_fields_on_open = True

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        reopened = docx.Document(stream)

        (field,) = reopened.fields
        assert field.type == "TOC"
        assert field.instruction == ' TOC \\o "1-3" \\h \\z \\u '
        assert field.dirty is True
        assert reopened.settings.update_fields_on_open is True


class DescribeInstructionBuilders:
    """Unit-test suite for the instruction builders in `docx.fields`."""

    @pytest.mark.parametrize(
        ("instruction", "expected"),
        [
            (fields.page_number(), " PAGE "),
            (fields.page_count(), " NUMPAGES "),
            (fields.table_of_contents(), ' TOC \\o "1-3" \\h \\z \\u '),
            (
                fields.table_of_contents(levels=(2, 4), hyperlinks=False),
                ' TOC \\o "2-4" \\z \\u ',
            ),
            (fields.cross_reference("intro"), " REF intro \\h "),
            (
                fields.cross_reference("intro", hyperlink=False, insert_paragraph_number=True),
                " REF intro \\n ",
            ),
            (fields.page_reference("intro"), " PAGEREF intro \\h "),
            (fields.sequence("Figure"), " SEQ Figure \\* ARABIC "),
            (
                fields.sequence("Figure", restart_at_heading_level=1),
                " SEQ Figure \\* ARABIC \\s 1 ",
            ),
            (fields.date(), " DATE "),
            (fields.date("d MMMM yyyy"), ' DATE \\@ "d MMMM yyyy" '),
            (fields.date(save_date=True), " SAVEDATE "),
            (fields.doc_property("Title"), " DOCPROPERTY Title "),
            (fields.doc_property("Project Name"), ' DOCPROPERTY "Project Name" '),
            (fields.styleref("Heading 1"), ' STYLEREF "Heading 1" '),
        ],
    )
    def it_builds_the_instruction(self, instruction: str, expected: str):
        assert instruction == expected

    @pytest.mark.parametrize("levels", [(0, 3), (1, 10), (3, 1)])
    def it_rejects_an_out_of_range_toc_level(self, levels: "tuple[int, int]"):
        with pytest.raises(ValueError, match="levels must be a range within 1-9"):
            fields.table_of_contents(levels=levels)

    def it_quotes_an_argument_containing_a_quote(self):
        assert fields.doc_property('a"b') == ' DOCPROPERTY "a""b" '

    def it_produces_instructions_that_read_back_as_the_right_field_type(self):
        """A round trip through the reader is the check that the syntax is well formed."""
        document = docx.Document()
        for instruction in (
            fields.page_number(),
            fields.table_of_contents(),
            fields.cross_reference("intro"),
            fields.sequence("Figure"),
            fields.doc_property("Title"),
        ):
            document.add_paragraph().add_field(instruction)

        assert [f.type for f in document.fields] == [
            "PAGE",
            "TOC",
            "REF",
            "SEQ",
            "DOCPROPERTY",
        ]


class DescribeUpdateFieldsSetting:
    """Unit-test suite for `docx.settings.Settings.update_fields_on_open`."""

    def it_is_off_by_default(self):
        assert docx.Document().settings.update_fields_on_open is False

    def it_can_be_turned_on_and_off(self):
        settings = docx.Document().settings

        settings.update_fields_on_open = True
        assert settings.update_fields_on_open is True

        settings.update_fields_on_open = False
        assert settings.update_fields_on_open is False
        assert settings._settings.xpath("./w:updateFields") == []
