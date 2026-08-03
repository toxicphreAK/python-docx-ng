# pyright: reportPrivateUsage=false

"""Unit test suite for the docx.bookmark module."""

from __future__ import annotations

import io

import pytest

import docx
from docx.bookmark import Bookmark, Bookmarks


class DescribeBookmarkCreation:
    """Unit-test suite for creating bookmarks."""

    def it_can_bookmark_a_paragraph(self):
        document = docx.Document()
        paragraph = document.add_paragraph("Introduction text")

        bookmark = paragraph.add_bookmark("Introduction")

        assert isinstance(bookmark, Bookmark)
        assert bookmark.name == "Introduction"
        assert bookmark.text == "Introduction text"
        assert bookmark.is_closed is True

    def it_places_the_delimiters_around_the_paragraph_content(self):
        document = docx.Document()
        paragraph = document.add_paragraph("body", style=None)
        paragraph.style = "Normal"

        paragraph.add_bookmark("Anchor")

        children = [child.tag.split("}")[1] for child in paragraph._p]
        # -- the start follows `w:pPr` and precedes the content; the end closes it --
        assert children == ["pPr", "bookmarkStart", "r", "bookmarkEnd"]

    def it_can_bookmark_a_range_of_runs(self):
        document = docx.Document()
        paragraph = document.add_paragraph()
        first = paragraph.add_run("start ")
        paragraph.add_run("middle ")
        last = paragraph.add_run("end")

        bookmark = first.mark_bookmark_range(last, "Range")

        assert bookmark.text == "start middle end"

    def it_can_bookmark_a_range_spanning_paragraphs(self):
        """A bookmark's delimiters are siblings of the content, not a container."""
        document = docx.Document()
        first = document.add_paragraph("first para ").runs[0]
        last = document.add_paragraph("second para").runs[0]

        bookmark = first.mark_bookmark_range(last, "Cross")

        assert bookmark.text == "first para \nsecond para"

    def it_gives_each_bookmark_a_distinct_id(self):
        document = docx.Document()

        ids = [document.add_paragraph("p").add_bookmark("B%d" % n).id for n in range(3)]

        assert ids == [0, 1, 2]
        assert len(set(ids)) == 3

    def it_survives_a_round_trip(self):
        document = docx.Document()
        document.add_paragraph("Introduction text").add_bookmark("Introduction")

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        reloaded = docx.Document(stream)
        assert reloaded.bookmarks["Introduction"].text == "Introduction text"


class DescribeBookmarks:
    """Unit-test suite for `docx.bookmark.Bookmarks`."""

    def it_provides_access_to_the_bookmarks_in_a_document(self, document: docx.Document):
        bookmarks = document.bookmarks

        assert isinstance(bookmarks, Bookmarks)
        assert [b.name for b in bookmarks] == ["First", "Second"]
        assert len(bookmarks) == 2

    def it_can_look_a_bookmark_up_by_name(self, document: docx.Document):
        assert document.bookmarks["Second"].text == "second"
        assert "First" in document.bookmarks
        assert "Absent" not in document.bookmarks

    def it_can_access_a_bookmark_by_index(self, document: docx.Document):
        assert document.bookmarks[0].name == "First"
        assert document.bookmarks[-1].name == "Second"

    def it_raises_looking_up_a_name_that_is_not_there(self, document: docx.Document):
        with pytest.raises(KeyError, match="no bookmark named"):
            document.bookmarks["Absent"]
        assert document.bookmarks.get("Absent") is None

    def it_leaves_out_the_bookmarks_word_maintains_for_itself(self):
        """`_GoBack` and the `_Toc…` anchors are not bookmarks the user made."""
        document = docx.Document()
        document.add_paragraph("mine").add_bookmark("Mine")
        document.add_paragraph("word's").add_bookmark("_GoBack")
        document.add_paragraph("toc").add_bookmark("_Toc12345")

        assert [b.name for b in document.bookmarks] == ["Mine"]
        assert [b.name for b in document.bookmarks.iter_all()] == [
            "Mine",
            "_GoBack",
            "_Toc12345",
        ]

    def it_can_delete_a_bookmark_leaving_its_content(self, document: docx.Document):
        document.bookmarks["First"].delete()

        assert [b.name for b in document.bookmarks] == ["Second"]
        assert document.paragraphs[0].text == "first"

    # fixtures ---------------------------------------------

    @pytest.fixture
    def document(self) -> docx.Document:
        document = docx.Document()
        document.add_paragraph("first").add_bookmark("First")
        document.add_paragraph("second").add_bookmark("Second")
        return document


class DescribeAnUnmatchedBookmark:
    """An unmatched delimiter is invalid but appears in real documents."""

    def it_reports_an_unclosed_bookmark_rather_than_raising(self):
        document = docx.Document()
        paragraph = document.add_paragraph("content")
        bookmark = paragraph.add_bookmark("Orphan")
        # -- remove the end delimiter, as a document from another tool might lack it --
        bookmarkEnd = bookmark._bookmarkStart.bookmarkEnd
        bookmarkEnd.getparent().remove(bookmarkEnd)

        assert bookmark.is_closed is False
        assert bookmark.text == ""
        assert [b.name for b in document.bookmarks] == ["Orphan"]

    def it_can_delete_an_unclosed_bookmark(self):
        document = docx.Document()
        bookmark = document.add_paragraph("content").add_bookmark("Orphan")
        bookmarkEnd = bookmark._bookmarkStart.bookmarkEnd
        bookmarkEnd.getparent().remove(bookmarkEnd)

        bookmark.delete()

        assert len(document.bookmarks) == 0
        assert document.paragraphs[0].text == "content"
