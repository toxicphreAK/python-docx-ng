# pyright: reportPrivateUsage=false

"""Unit test suite for the docx.altchunk module."""

from __future__ import annotations

import io
from typing import cast

import pytest

import docx
from docx.altchunk import AltChunk
from docx.oxml.document import CT_AltChunk
from docx.parts.altchunk import AltChunkPart
from docx.parts.document import DocumentPart

from .unitutil.cxml import element
from .unitutil.mock import FixtureRequest, Mock, instance_mock


class DescribeAltChunk:
    """Unit-test suite for `docx.altchunk.AltChunk`."""

    def it_provides_access_to_the_embedded_document(
        self, document_part_: Mock, alt_chunk_part_: Mock
    ):
        alt_chunk_part_.blob = b"<html/>"
        alt_chunk_part_.content_type = "text/html"
        document_part_.part = document_part_
        document_part_.related_parts = {"rId6": alt_chunk_part_}
        altChunk = cast(CT_AltChunk, element("w:altChunk{r:id=rId6}"))

        alt_chunk = AltChunk(altChunk, document_part_)

        assert alt_chunk.blob == b"<html/>"
        assert alt_chunk.content_type == "text/html"

    def but_it_raises_when_it_names_no_content(self, document_part_: Mock):
        alt_chunk = AltChunk(cast(CT_AltChunk, element("w:altChunk")), document_part_)

        with pytest.raises(ValueError, match="has no `r:id`"):
            alt_chunk.blob

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def alt_chunk_part_(self, request: FixtureRequest):
        return instance_mock(request, AltChunkPart)

    @pytest.fixture
    def document_part_(self, request: FixtureRequest):
        return instance_mock(request, DocumentPart)


class DescribeAltChunkRoundTrip:
    """Integration-test suite for alt-chunk persistence."""

    def it_survives_a_save_and_reload(self):
        document = docx.Document()
        document.add_paragraph("before")
        document.add_alt_chunk(b"<html><body><p>imported</p></body></html>", "text/html")
        document.add_alt_chunk(io.BytesIO(rb"{\rtf1}"), "application/rtf")

        saved = io.BytesIO()
        document.save(saved)
        reloaded = docx.Document(io.BytesIO(saved.getvalue()))

        assert [(a.content_type, a.blob) for a in reloaded.alt_chunks] == [
            ("text/html", b"<html><body><p>imported</p></body></html>"),
            ("application/rtf", rb"{\rtf1}"),
        ]

    def and_its_content_stays_invisible_to_the_document_body(self):
        document = docx.Document()
        document.add_paragraph("only paragraph")
        document.add_alt_chunk(b"<html><body><p>imported</p></body></html>", "text/html")

        saved = io.BytesIO()
        document.save(saved)
        reloaded = docx.Document(io.BytesIO(saved.getvalue()))

        # -- Word does the import when it opens the file; until then the embedded
        # -- content is opaque to this library --
        assert [p.text for p in reloaded.paragraphs] == ["only paragraph"]
        assert list(reloaded.iter_inner_content()) != []
        assert len(reloaded.alt_chunks) == 1
