# pyright: reportPrivateUsage=false

"""Unit test suite for the caption API — issue #102."""

from __future__ import annotations

import io

import docx
from docx import fields
from docx.caption import Caption


def _reopened(document):
    stream = io.BytesIO()
    document.save(stream)
    return docx.Document(io.BytesIO(stream.getvalue()))


class DescribeAddCaption:
    """Unit-test suite for `Document.add_caption()`."""

    def it_returns_a_caption(self):
        document = docx.Document()

        caption = document.add_caption("Figure", "Cross-section")

        assert isinstance(caption, Caption)
        assert caption.label == "Figure"
        assert caption.paragraph._p is document.paragraphs[0]._p

    def it_writes_the_label_a_SEQ_field_and_the_text(self):
        document = docx.Document()

        caption = document.add_caption("Figure", "Cross-section")

        seq = [f for f in caption.paragraph.fields if f.type == "SEQ"]
        assert len(seq) == 1
        assert "Figure" in seq[0].instruction
        assert "\\* ARABIC" in seq[0].instruction
        assert caption.text.startswith("Figure")
        assert caption.text.endswith("Cross-section")

    def it_applies_the_Caption_style(self):
        document = docx.Document()

        caption = document.add_caption("Figure", "x")

        assert caption.paragraph.style.name == "Caption"

    def and_the_style_can_be_changed_or_skipped(self):
        document = docx.Document()

        styled = document.add_caption("Figure", "x", style="Quote")
        plain = document.add_caption("Figure", "y", style=None)

        assert styled.paragraph.style.name == "Quote"
        assert plain.paragraph.style.name == "Normal"

    def it_can_omit_the_caption_text(self):
        document = docx.Document()

        caption = document.add_caption("Figure")

        assert caption.text.strip() == "Figure"

    def it_can_use_a_different_separator(self):
        document = docx.Document()

        caption = document.add_caption("Figure", "x", separator=": ")

        assert caption.text.endswith(": x")

    def it_can_restart_numbering_at_a_heading_level(self):
        """The "Figure 2-1" style, which the `\\s` switch gives."""
        document = docx.Document()

        caption = document.add_caption("Figure", "x", restart_at_heading_level=1)

        seq = [f for f in caption.paragraph.fields if f.type == "SEQ"][0]
        assert "\\s 1" in seq.instruction

    def it_can_be_placed_before_an_existing_paragraph(self):
        """Which is where a table caption goes."""
        document = docx.Document()
        target = document.add_paragraph("body")

        document.add_caption("Table", "x", before=target)

        assert [p.text for p in document.paragraphs][1] == "body"
        assert document.paragraphs[0].style.name == "Caption"

    def it_reports_no_number_until_Word_computes_one(self):
        """A `SEQ` field's result is cached by Word; a fresh field has none."""
        document = docx.Document()

        caption = document.add_caption("Figure", "x")

        assert caption.number is None

    def it_survives_a_save(self):
        document = docx.Document()
        document.add_caption("Figure", "Cross-section")

        reopened = _reopened(document)

        assert reopened.paragraphs[0].style.name == "Caption"
        assert [f.type for f in reopened.paragraphs[0].fields] == ["SEQ"]

    def it_has_a_useful_repr(self):
        caption = docx.Document().add_caption("Figure", "x")

        assert "Caption" in repr(caption)
        assert "_Ref" in repr(caption)


class DescribeCaptionBookmarking:
    """The `_Ref` naming is what makes a caption referenceable from Word's UI."""

    def it_bookmarks_the_caption_with_a_Ref_prefixed_name(self):
        document = docx.Document()

        caption = document.add_caption("Figure", "x")

        assert caption.bookmark_name.startswith("_Ref")
        names = document._element.xpath("//w:bookmarkStart/@w:name")
        assert caption.bookmark_name in names

    def and_each_caption_gets_a_distinct_name(self):
        document = docx.Document()

        first = document.add_caption("Figure", "one")
        second = document.add_caption("Figure", "two")
        third = document.add_caption("Table", "three")

        assert len({first.bookmark_name, second.bookmark_name, third.bookmark_name}) == 3

    def and_the_names_do_not_collide_with_existing_ones(self):
        document = docx.Document()
        document.add_paragraph("x").add_bookmark("_Ref000000005")

        caption = document.add_caption("Figure", "x")

        assert caption.bookmark_name == "_Ref000000006"

    def the_bookmark_is_hidden_from_the_ordinary_collection(self):
        """`Document.bookmarks` leaves out the ones Word maintains for itself."""
        document = docx.Document()

        caption = document.add_caption("Figure", "x")

        assert [b.name for b in document.bookmarks] == []
        assert caption.bookmark_name in [b.name for b in document.bookmarks.iter_all()]

    def it_makes_a_cross_reference_a_one_liner(self):
        document = docx.Document()
        caption = document.add_caption("Figure", "Cross-section")

        paragraph = document.add_paragraph("see ")
        paragraph.add_field(fields.cross_reference(caption.bookmark_name))

        ref = [f for f in paragraph.fields if f.type == "REF"][0]
        assert caption.bookmark_name in ref.instruction

    def and_the_reference_resolves_after_a_save(self):
        document = docx.Document()
        caption = document.add_caption("Figure", "Cross-section")
        document.add_paragraph().add_field(fields.cross_reference(caption.bookmark_name))

        reopened = _reopened(document)

        names = reopened._element.xpath("//w:bookmarkStart/@w:name")
        ref = [f for f in reopened.paragraphs[1].fields if f.type == "REF"][0]
        assert any(name in ref.instruction for name in names)

    def the_names_are_stable_across_runs(self):
        """A timestamp-derived name would break byte-reproducible output."""
        first = docx.Document().add_caption("Figure", "x").bookmark_name
        second = docx.Document().add_caption("Figure", "x").bookmark_name

        assert first == second


class DescribeCellAddCaption:
    """Unit-test suite for `_Cell.add_caption()`."""

    def it_adds_a_caption_to_a_cell(self):
        document = docx.Document()
        cell = document.add_table(1, 1).cell(0, 0)

        caption = cell.add_caption("Table", "in a cell")

        assert caption.paragraph._p in [p._p for p in cell.paragraphs]
        assert caption.paragraph.style.name == "Caption"
        assert caption.bookmark_name.startswith("_Ref")
