"""Unit test suite for the tracked-changes API."""

from __future__ import annotations

import datetime as dt
import io

import pytest

import docx
from docx.enum.revision import WD_REVISION_TYPE
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import parse_xml

_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _document(body_inner_xml: str) -> docx.document.Document:
    """A document whose body is `body_inner_xml`, keeping the template's `w:sectPr`."""
    document = docx.Document()
    body = parse_xml(f"<w:body {nsdecls('w')}>{body_inner_xml}</w:body>")
    for child in list(document.element.body):
        if child.tag == qn("w:sectPr"):
            body.append(child)
    document.element.replace(document.element.body, body)
    return document


_REVISED_PARAGRAPH = """
<w:p>
  <w:r><w:t xml:space="preserve">The </w:t></w:r>
  <w:ins w:id="1" w:author="Ada" w:date="2026-01-02T10:00:00Z">
    <w:r><w:t xml:space="preserve">quick </w:t></w:r>
  </w:ins>
  <w:del w:id="2" w:author="Bob" w:date="2026-01-03T11:30:00Z">
    <w:r><w:delText xml:space="preserve">slow </w:delText></w:r>
  </w:del>
  <w:r><w:t>fox</w:t></w:r>
</w:p>
"""

_SPLIT_PARAGRAPHS = """
<w:p><w:pPr><w:rPr><w:{tag} w:id="9" w:author="Ada"/></w:rPr></w:pPr>
  <w:r><w:t xml:space="preserve">first half </w:t></w:r></w:p>
<w:p><w:r><w:t>second half</w:t></w:r></w:p>
"""

_FORMATTING_CHANGE = """
<w:p><w:r>
  <w:rPr><w:b/>
    <w:rPrChange w:id="4" w:author="Ada"><w:rPr><w:i/></w:rPr></w:rPrChange>
  </w:rPr>
  <w:t>styled</w:t>
</w:r></w:p>
"""

_ROW_REVISION = """
<w:tbl>
  <w:tr><w:tc><w:p><w:r><w:t>kept</w:t></w:r></w:p></w:tc></w:tr>
  <w:tr><w:trPr><w:{tag} w:id="3" w:author="Ada"/></w:trPr>
    <w:tc><w:p><w:r><w:t>marked</w:t></w:r></w:p></w:tc></w:tr>
</w:tbl>
"""


class DescribeRevisionTextModel:
    """Unit-test suite for the two defined readings of a revised paragraph."""

    def it_reads_the_document_as_it_now_stands(self):
        """Insertions in, deletions out — the document with every revision accepted."""
        document = _document(_REVISED_PARAGRAPH)

        assert document.paragraphs[0].text == "The quick fox"

    def it_reads_the_document_as_it_stood_before(self):
        document = _document(_REVISED_PARAGRAPH)

        assert document.paragraphs[0].original_text == "The slow fox"

    def it_reads_both_the_same_for_a_paragraph_with_no_revisions(self):
        document = _document("<w:p><w:r><w:t>plain</w:t></w:r></w:p>")
        paragraph = document.paragraphs[0]

        assert paragraph.text == paragraph.original_text == "plain"

    def it_includes_inserted_runs_in_the_run_collection(self):
        document = _document(_REVISED_PARAGRAPH)

        texts = [r.text for r in document.paragraphs[0].runs]

        assert texts == ["The ", "quick ", "fox"], "the deleted run is not in the text"

    def it_reads_text_moved_to_here_as_present(self):
        document = _document(
            '<w:p><w:moveTo w:id="1" w:author="Ada"><w:r><w:t>moved</w:t></w:r></w:moveTo></w:p>'
        )
        paragraph = document.paragraphs[0]

        assert paragraph.text == "moved"
        assert paragraph.original_text == ""

    def it_reads_text_moved_away_as_absent(self):
        document = _document(
            '<w:p><w:moveFrom w:id="1" w:author="Ada">'
            "<w:r><w:delText>gone</w:delText></w:r></w:moveFrom></w:p>"
        )
        paragraph = document.paragraphs[0]

        assert paragraph.text == ""
        assert paragraph.original_text == "gone"

    def it_reads_a_revision_inside_a_hyperlink(self):
        document = _document(
            '<w:p><w:hyperlink r:id="rId9"><w:ins w:id="1" w:author="Ada">'
            "<w:r><w:t>linked</w:t></w:r></w:ins></w:hyperlink></w:p>".replace(
                "<w:hyperlink", f'<w:hyperlink xmlns:r="{_R_NS}"', 1
            )
        )
        paragraph = document.paragraphs[0]

        assert paragraph.text == "linked"
        assert paragraph.original_text == ""


class DescribeReadingRevisions:
    """Unit-test suite for `Document.revisions` and `Paragraph.revisions`."""

    def it_reports_each_revision_with_its_metadata(self):
        document = _document(_REVISED_PARAGRAPH)

        insertion, deletion = document.revisions

        assert insertion.type == WD_REVISION_TYPE.INSERTION
        assert insertion.author == "Ada"
        assert insertion.id == 1
        assert insertion.text == "quick "
        assert insertion.date == dt.datetime(2026, 1, 2, 10, 0, tzinfo=dt.timezone.utc)

        assert deletion.type == WD_REVISION_TYPE.DELETION
        assert deletion.author == "Bob"
        assert deletion.text == "slow ", "the point of a deletion is what it removed"

    def it_reports_revisions_per_paragraph_too(self):
        document = _document(_REVISED_PARAGRAPH)

        assert len(document.paragraphs[0].revisions) == 2

    def it_is_empty_for_a_document_that_has_not_been_reviewed(self):
        document = _document("<w:p><w:r><w:t>plain</w:t></w:r></w:p>")

        assert document.revisions == []

    @pytest.mark.parametrize(
        ("date_attr", "expected"),
        [
            (
                'w:date="2026-01-02T10:00:00Z"',
                dt.datetime(2026, 1, 2, 10, 0, tzinfo=dt.timezone.utc),
            ),
            ("", None),
            ('w:date=""', None),
            ('w:date="not a date"', None),
        ],
    )
    def it_tolerates_a_missing_or_unparseable_date(self, date_attr: str, expected):
        """An anonymised document has the timestamp stripped or blanked."""
        document = _document(
            f'<w:p><w:ins w:id="1" w:author="Ada" {date_attr}><w:r><w:t>x</w:t></w:r></w:ins></w:p>'
        )

        assert document.revisions[0].date == expected

    def it_tolerates_a_revision_with_no_author(self):
        document = _document('<w:p><w:ins w:id="1"><w:r><w:t>x</w:t></w:r></w:ins></w:p>')

        assert document.revisions[0].author == ""

    def it_knows_a_paragraph_mark_revision(self):
        document = _document(_SPLIT_PARAGRAPHS.format(tag="del"))

        (revision,) = document.revisions

        assert revision.is_paragraph_mark is True
        assert revision.is_row is False
        assert revision.text == ""

    def it_knows_a_row_revision(self):
        document = _document(_ROW_REVISION.format(tag="ins"))

        (revision,) = document.revisions

        assert revision.is_row is True
        assert revision.type == WD_REVISION_TYPE.INSERTION

    def it_knows_a_formatting_change(self):
        document = _document(_FORMATTING_CHANGE)

        (revision,) = document.revisions

        assert revision.type == WD_REVISION_TYPE.FORMATTING
        assert revision.author == "Ada"


class DescribeAcceptAndReject:
    """Unit-test suite for applying revisions."""

    def it_accepts_content_revisions(self):
        document = _document(_REVISED_PARAGRAPH)

        applied = document.accept_all_revisions()

        assert applied == 2
        assert document.paragraphs[0].text == "The quick fox"
        assert document.revisions == [], "the records are gone"
        assert document.element.xpath(".//w:delText") == []

    def it_rejects_content_revisions(self):
        document = _document(_REVISED_PARAGRAPH)

        applied = document.reject_all_revisions()

        assert applied == 2
        assert document.paragraphs[0].text == "The slow fox"
        assert document.revisions == []

    def it_turns_restored_deleted_text_back_into_ordinary_text(self):
        """`w:delText` is how a deletion hides text; restoring it must undo that."""
        document = _document(_REVISED_PARAGRAPH)

        document.reject_all_revisions()

        assert document.element.xpath(".//w:delText") == []
        assert "slow " in [str(t) for t in document.element.xpath(".//w:t")]

    def it_preserves_significant_whitespace_when_restoring_deleted_text(self):
        document = _document(_REVISED_PARAGRAPH)

        document.reject_all_revisions()

        restored = next(t for t in document.element.xpath(".//w:t") if str(t) == "slow ")
        assert restored.get(qn("xml:space")) == "preserve"

    def it_can_apply_a_single_revision(self):
        document = _document(_REVISED_PARAGRAPH)

        document.revisions[0].accept()

        assert document.paragraphs[0].text == "The quick fox"
        assert len(document.revisions) == 1, "the deletion is still pending"

    def it_can_apply_revisions_per_paragraph(self):
        document = _document(_REVISED_PARAGRAPH)

        applied = document.paragraphs[0].accept_all_revisions()

        assert applied == 2
        assert document.paragraphs[0].revisions == []

    @pytest.mark.parametrize(
        ("tag", "accept", "expected"),
        [
            # -- a deleted paragraph mark means the two paragraphs are really one --
            ("del", True, ["first half second half"]),
            ("del", False, ["first half ", "second half"]),
            # -- an inserted one means the split is the change --
            ("ins", True, ["first half ", "second half"]),
            ("ins", False, ["first half second half"]),
        ],
    )
    def it_merges_or_splits_paragraphs_for_a_paragraph_mark_revision(
        self, tag: str, accept: bool, expected: list[str]
    ):
        document = _document(_SPLIT_PARAGRAPHS.format(tag=tag))

        document.accept_all_revisions() if accept else document.reject_all_revisions()

        assert [p.text for p in document.paragraphs] == expected
        assert document.revisions == []

    @pytest.mark.parametrize(
        ("tag", "accept", "expected_rows"),
        [
            ("ins", True, 2),
            ("ins", False, 1),
            ("del", True, 1),
            ("del", False, 2),
        ],
    )
    def it_adds_or_removes_a_row_for_a_row_revision(
        self, tag: str, accept: bool, expected_rows: int
    ):
        document = _document(_ROW_REVISION.format(tag=tag))

        document.accept_all_revisions() if accept else document.reject_all_revisions()

        assert len(document.tables[0].rows) == expected_rows

    def it_drops_the_record_when_accepting_a_formatting_change(self):
        document = _document(_FORMATTING_CHANGE)
        run = document.paragraphs[0].runs[0]
        assert (run.bold, run.italic) == (True, None)

        document.accept_all_revisions()

        run = document.paragraphs[0].runs[0]
        assert (run.bold, run.italic) == (True, None), "the current formatting stands"
        assert document.element.xpath(".//w:rPrChange") == []

    def it_restores_the_previous_formatting_when_rejecting(self):
        document = _document(_FORMATTING_CHANGE)

        document.reject_all_revisions()

        run = document.paragraphs[0].runs[0]
        assert (run.bold, run.italic) == (None, True), "the recorded formatting is back"
        assert document.element.xpath(".//w:rPrChange") == []

    def it_applies_a_nested_revision_before_the_one_containing_it(self):
        """Unwrapping the outer revision first would leave the inner one detached."""
        document = _document(
            '<w:p><w:ins w:id="1" w:author="Ada">'
            "<w:r><w:t>outer </w:t></w:r>"
            '<w:del w:id="2" w:author="Bob"><w:r><w:delText>inner</w:delText></w:r></w:del>'
            "</w:ins></w:p>"
        )

        applied = document.accept_all_revisions()

        assert applied == 2
        assert document.paragraphs[0].text == "outer "
        assert document.revisions == []

    def it_survives_a_round_trip(self):
        document = _document(_REVISED_PARAGRAPH)
        document.accept_all_revisions()

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        reopened = docx.Document(stream)

        assert reopened.paragraphs[0].text == "The quick fox"
        assert reopened.revisions == []


class DescribeTrackRevisionsSetting:
    """Unit-test suite for `docx.settings.Settings.track_revisions`."""

    def it_is_off_by_default(self):
        assert docx.Document().settings.track_revisions is False

    def it_can_be_turned_on_and_off(self):
        settings = docx.Document().settings

        settings.track_revisions = True
        assert settings.track_revisions is True

        settings.track_revisions = False
        assert settings.track_revisions is False
        assert settings._settings.xpath("./w:trackRevisions") == []

    def it_survives_a_round_trip(self):
        document = docx.Document()
        document.settings.track_revisions = True

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        assert docx.Document(stream).settings.track_revisions is True
