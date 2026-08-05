"""Every entry point that takes a path takes an `os.PathLike` as well as a `str`.

The README states this without qualification, so it is a claim the tests have to hold
up. Opening and saving already accepted one; the image, alt-chunk and embedded-object
paths reached `open()` through helpers that tested `isinstance(x, str)` and fell through
to the stream branch, where a `Path` fails with `AttributeError: no attribute 'seek'` —
a message that says nothing about what was wrong.
"""

from __future__ import annotations

import pathlib

import pytest

import docx
from docx.image.image import Image

from .unitutil.file import test_file

_IMAGE = pathlib.Path(test_file("monty-truth.png"))


class DescribeImagePathAcceptance:
    """The image loader is the one every picture entry point funnels through."""

    def it_loads_an_image_from_a_path_object(self):
        image = Image.from_file(_IMAGE)

        assert (image.px_width, image.px_height) == (150, 214)
        # -- the filename survives, which is what the extension is taken from --
        assert image.ext == "png"

    def and_a_str_path_still_works(self):
        assert Image.from_file(str(_IMAGE)).ext == "png"


class DescribePathLikeEntryPoints:
    """Each of these reached `open()` through a helper that only knew about `str`."""

    def it_accepts_a_path_for_add_picture(self):
        document = docx.Document()

        shape = document.add_picture(_IMAGE)

        assert shape.width is not None

    def and_for_add_float_picture(self):
        run = docx.Document().add_paragraph().add_run()

        shape = run.add_float_picture(_IMAGE)

        assert shape.width is not None

    def and_for_add_image_watermark(self):
        document = docx.Document()

        document.add_image_watermark(_IMAGE)

        assert len(document.watermarks) > 0

    def and_for_add_alt_chunk(self):
        document = docx.Document()

        chunk = document.add_alt_chunk(_IMAGE, "text/html")

        assert chunk.blob == _IMAGE.read_bytes()

    def and_for_add_embedded_object(self):
        run = docx.Document().add_paragraph().add_run()

        obj = run.add_embedded_object(_IMAGE, icon=_IMAGE)

        assert obj.blob is not None

    @pytest.mark.parametrize("as_str", [True, False])
    def it_gives_the_same_result_either_way(self, as_str: bool):
        """A `Path` and its `str` are the same path, so they are the same document."""
        document = docx.Document()

        document.add_picture(str(_IMAGE) if as_str else _IMAGE)

        assert len(document.images) == 1
