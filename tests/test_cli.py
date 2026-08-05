# pyright: reportPrivateUsage=false

"""Unit test suite for the `python -m docx` CLI — issue #118."""

from __future__ import annotations

import io
import json

import docx
from docx.cli import EXIT_CANNOT_OPEN, EXIT_WOULD_CHANGE, main


def _run(*argv: str) -> tuple[int, str]:
    """Run the CLI with `argv` and return `(exit_code, stdout)`."""
    out = io.StringIO()
    code = main(argv, stdout=out)
    return code, out.getvalue()


def _document_path(tmp_path, name: str = "doc.docx"):
    document = docx.Document()
    document.add_heading("H", 1)
    document.add_paragraph("body")
    path = tmp_path / name
    document.save(str(path))
    return str(path)


class DescribeInfoCommand:
    """`info` is the fastest answer to "why is this file like this"."""

    def it_lists_the_parts_with_their_sizes(self, tmp_path):
        path = _document_path(tmp_path)

        code, out = _run("info", path)

        assert code == 0
        assert "word/styles.xml" in out
        assert "TOTAL uncompressed" in out

    def it_sorts_the_biggest_part_first(self, tmp_path):
        path = _document_path(tmp_path)

        _, out = _run("info", path)

        first_line = out.splitlines()[0]
        assert "word/styles.xml" in first_line

    def it_can_emit_json(self, tmp_path):
        path = _document_path(tmp_path)

        _, out = _run("info", "--json", path)

        payload = json.loads(out)
        assert payload["uncompressed"] > 0
        assert any(part["name"] == "word/styles.xml" for part in payload["parts"])

    def it_exits_non_zero_for_a_file_that_is_not_there(self, tmp_path):
        code, _ = _run("info", str(tmp_path / "nope.docx"))

        assert code == EXIT_CANNOT_OPEN

    def and_for_a_file_that_is_not_a_package(self, tmp_path):
        path = tmp_path / "not.docx"
        path.write_bytes(b"not a zip")

        code, _ = _run("info", str(path))

        assert code == EXIT_CANNOT_OPEN


class DescribeStylesCommands:
    """The `styles` subcommands are thin shells over `Styles.usage()`."""

    def it_reports_defined_and_used_counts(self, tmp_path):
        path = _document_path(tmp_path)

        code, out = _run("styles", "report", path)

        assert code == 0
        assert "styles defined" in out
        assert "in use" in out
        assert "latent style exceptions" in out

    def and_as_json(self, tmp_path):
        path = _document_path(tmp_path)

        _, out = _run("styles", "report", "--json", path)

        payload = json.loads(out)
        assert payload["defined"] > payload["used"]
        assert payload["unused"] == payload["defined"] - payload["used"]
        assert payload["by_type"]["paragraph"] > 0

    def it_lists_every_style_by_default(self, tmp_path):
        path = _document_path(tmp_path)

        _, out = _run("styles", "list", path)

        assert "Heading 1" in out
        assert "Quote" in out

    def and_only_the_used_ones_when_asked(self, tmp_path):
        path = _document_path(tmp_path)

        _, out = _run("styles", "list", "--used", path)

        assert "Heading 1" in out
        assert "Quote" not in out

    def and_only_the_unused_ones(self, tmp_path):
        path = _document_path(tmp_path)

        _, out = _run("styles", "list", "--unused", path)

        assert "Quote" in out
        assert "Heading 1\n" not in out

    def it_extracts_named_styles_to_a_new_document(self, tmp_path):
        path = _document_path(tmp_path)
        output = str(tmp_path / "styles-only.docx")

        code, out = _run(
            "styles", "extract", path, "-o", output, "--names", "Quote,Caption"
        )

        assert code == 0
        assert "extracted" in out
        extracted = docx.Document(output)
        assert "Quote" in extracted.styles
        assert "Caption" in extracted.styles
        assert len(extracted.paragraphs) == 0

    def and_it_can_write_a_template(self, tmp_path):
        from docx.opc.constants import CONTENT_TYPE as CT

        path = _document_path(tmp_path)
        output = str(tmp_path / "house.dotx")

        _run("styles", "extract", path, "-o", output, "--names", "Quote", "--as-template")

        assert docx.Document(output).part.content_type == CT.WML_TEMPLATE_MAIN

    def the_bare_styles_command_prints_help(self, tmp_path):
        code, out = _run("styles")

        assert code == 0
        assert "report" in out


class DescribeCleanupCommand:
    """Destructive, so the rules about writing and exit codes are the point."""

    def it_never_writes_to_the_input(self, tmp_path):
        """Someone will point this at their only copy."""
        path = _document_path(tmp_path)
        with open(path, "rb") as f:
            before = f.read()

        code, _ = _run("cleanup", path)

        assert code == EXIT_CANNOT_OPEN
        with open(path, "rb") as f:
            assert f.read() == before

    def it_writes_a_smaller_document_to_the_output(self, tmp_path):
        path = _document_path(tmp_path)
        output = str(tmp_path / "clean.docx")

        code, out = _run("cleanup", path, "-o", output, "--latent-styles")

        assert code == 0
        assert "wrote" in out
        import os

        assert os.path.getsize(output) < os.path.getsize(path)

    def and_the_result_still_reads(self, tmp_path):
        path = _document_path(tmp_path)
        output = str(tmp_path / "clean.docx")

        _run("cleanup", path, "-o", output)

        cleaned = docx.Document(output)
        assert [p.text for p in cleaned.paragraphs] == ["H", "body"]
        assert cleaned.paragraphs[0].style.name == "Heading 1"

    def check_reports_without_writing_and_exits_non_zero(self, tmp_path):
        path = _document_path(tmp_path)

        code, out = _run("cleanup", "--check", path)

        assert code == EXIT_WOULD_CHANGE
        assert "would have removed" in out
        assert not (tmp_path / "clean.docx").exists()

    def and_exits_zero_when_there_is_nothing_to_remove(self, tmp_path):
        """So it can be a CI gate."""
        path = _document_path(tmp_path)
        output = str(tmp_path / "clean.docx")
        _run("cleanup", path, "-o", output, "--latent-styles")

        code, _ = _run("cleanup", "--check", output, "--latent-styles")

        assert code == 0

    def it_keeps_the_styles_named_in_keep(self, tmp_path):
        path = _document_path(tmp_path)
        output = str(tmp_path / "clean.docx")

        _run("cleanup", path, "-o", output, "--keep", "Quote,Caption")

        cleaned = docx.Document(output)
        assert "Quote" in cleaned.styles
        assert "Caption" in cleaned.styles

    def each_pass_can_be_turned_off(self, tmp_path):
        path = _document_path(tmp_path)
        output = str(tmp_path / "clean.docx")

        _, out = _run(
            "cleanup",
            path,
            "-o",
            output,
            "--no-styles",
            "--no-numbering",
            "--no-media",
            "--json",
        )

        payload = json.loads(out)
        assert payload["styles"] == []
        assert payload["num_ids"] == []
        assert payload["media"] == []

    def it_leaves_latent_styles_alone_unless_asked(self, tmp_path):
        path = _document_path(tmp_path)
        output = str(tmp_path / "clean.docx")

        _, out = _run("cleanup", path, "-o", output, "--json")

        assert json.loads(out)["latent_styles"] == 0


class DescribeBareInvocation:
    def it_prints_help_with_no_arguments(self):
        code, out = _run()

        assert code == 0
        assert "Inspect and clean up Word documents" in out
