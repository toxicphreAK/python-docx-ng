"""Test suite for the docx.api module."""

import io
import re
import zipfile

import pytest

from docx.api import Document as DocumentFactoryFn
from docx.document import Document as DocumentCls
from docx.opc.constants import CONTENT_TYPE as CT

from .unitutil.file import test_file
from .unitutil.mock import FixtureRequest, Mock, class_mock, function_mock, instance_mock


class DescribeDocument:
    """Unit-test suite for `docx.api.Document` factory function."""

    def it_opens_a_docx_file(self, Package_: Mock, document_: Mock):
        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = CT.WML_DOCUMENT_MAIN

        document = DocumentFactoryFn("foobar.docx")

        Package_.open.assert_called_once_with("foobar.docx")
        assert document is document_

    def it_opens_the_default_docx_if_none_specified(
        self, _default_docx_path_: Mock, Package_: Mock, document_: Mock
    ):
        _default_docx_path_.return_value = "default-document.docx"
        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = CT.WML_DOCUMENT_MAIN

        document = DocumentFactoryFn()

        Package_.open.assert_called_once_with("default-document.docx")
        assert document is document_

    def it_opens_a_docx_file_from_a_pathlib_path(self, Package_: Mock, document_: Mock):
        from pathlib import Path

        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = CT.WML_DOCUMENT_MAIN

        document = DocumentFactoryFn(Path("foobar.docx"))

        Package_.open.assert_called_once_with("foobar.docx")
        assert document is document_

    def it_opens_a_macro_enabled_docm_file(self, Package_: Mock, document_: Mock):
        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = CT.WML_DOCUMENT_MACRO_ENABLED_MAIN

        document = DocumentFactoryFn("foobar.docm")

        Package_.open.assert_called_once_with("foobar.docm")
        assert document is document_

    @pytest.mark.parametrize(
        ("filename", "content_type"),
        [
            ("foobar.dotx", CT.WML_TEMPLATE_MAIN),
            ("foobar.dotm", CT.WML_TEMPLATE_MACRO_ENABLED_MAIN),
        ],
    )
    def it_opens_a_word_template(
        self, filename: str, content_type: str, Package_: Mock, document_: Mock
    ):
        document_part = Package_.open.return_value.main_document_part
        document_part.document = document_
        document_part.content_type = content_type

        document = DocumentFactoryFn(filename)

        Package_.open.assert_called_once_with(filename)
        assert document is document_

    def it_raises_on_not_a_Word_file(self, Package_: Mock):
        Package_.open.return_value.main_document_part.content_type = "BOGUS"

        with pytest.raises(ValueError, match="file 'foobar.xlsx' is not a Word file,"):
            DocumentFactoryFn("foobar.xlsx")

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def _default_docx_path_(self, request: FixtureRequest):
        return function_mock(request, "docx.api._default_docx_path")

    @pytest.fixture
    def document_(self, request: FixtureRequest):
        return instance_mock(request, DocumentCls)

    @pytest.fixture
    def Package_(self, request: FixtureRequest):
        return class_mock(request, "docx.api.Package")


class DescribeMacroEnabledDocuments:
    """Integration-test suite for `.docm` support."""

    def it_can_open_and_edit_a_docm_file(self):
        document = DocumentFactoryFn(test_file("macro-enabled.docm"))

        document.add_paragraph("Hello")

        assert document.paragraphs[-1].text == "Hello"

    def it_preserves_the_macro_storage_when_saving(self):
        """The vbaProject part has no API but must survive a round trip intact."""
        document = DocumentFactoryFn(test_file("macro-enabled.docm"))
        with zipfile.ZipFile(test_file("macro-enabled.docm")) as z:
            original_vba = z.read("word/vbaProject.bin")

        stream = io.BytesIO()
        document.save(stream)

        with zipfile.ZipFile(stream) as z:
            assert z.read("word/vbaProject.bin") == original_vba

    def it_keeps_the_macro_enabled_content_type_when_saving(self):
        """Re-saving as a plain document content type would break macros in Word."""
        document = DocumentFactoryFn(test_file("macro-enabled.docm"))

        stream = io.BytesIO()
        document.save(stream)

        with zipfile.ZipFile(stream) as z:
            content_types = z.read("[Content_Types].xml").decode("utf-8")
        assert CT.WML_DOCUMENT_MACRO_ENABLED_MAIN in content_types


class DescribeWordTemplates:
    """Integration-test suite for `.dotx` and `.dotm` support."""

    def it_can_open_and_edit_a_dotx_template(self):
        document = DocumentFactoryFn(test_file("template.dotx"))

        document.add_paragraph("Hello")

        assert document.is_template is True
        assert document.paragraphs[-1].text == "Hello"

    def it_can_open_a_macro_enabled_dotm_template(self):
        document = DocumentFactoryFn(test_file("macro-enabled.dotm"))

        assert document.is_template is True

    def it_stays_a_template_when_saved(self):
        """Opening and saving a template must not silently turn it into a document."""
        document = DocumentFactoryFn(test_file("template.dotx"))

        stream = io.BytesIO()
        document.save(stream)

        assert _main_part_content_type(stream) == CT.WML_TEMPLATE_MAIN

    def it_can_save_a_template_as_an_ordinary_document(self):
        """Generating a document from a template is the point of opening one."""
        document = DocumentFactoryFn(test_file("template.dotx"))

        stream = io.BytesIO()
        document.save(stream, as_template=False)

        assert _main_part_content_type(stream) == CT.WML_DOCUMENT_MAIN
        assert document.is_template is False
        # -- and the result opens as a document --
        stream.seek(0)
        assert DocumentFactoryFn(stream).is_template is False

    def it_keeps_a_macro_enabled_template_macro_enabled_when_converting(self):
        document = DocumentFactoryFn(test_file("macro-enabled.dotm"))

        stream = io.BytesIO()
        document.save(stream, as_template=False)

        assert _main_part_content_type(stream) == CT.WML_DOCUMENT_MACRO_ENABLED_MAIN

    def it_can_save_an_ordinary_document_as_a_template(self):
        document = DocumentFactoryFn(test_file("test.docx"))

        stream = io.BytesIO()
        document.save(stream, as_template=True)

        assert _main_part_content_type(stream) == CT.WML_TEMPLATE_MAIN

    def it_knows_an_ordinary_document_is_not_a_template(self):
        assert DocumentFactoryFn(test_file("test.docx")).is_template is False


def _main_part_content_type(stream: io.BytesIO) -> str:
    """The content-type override declared for `/word/document.xml` in `stream`."""
    stream.seek(0)
    with zipfile.ZipFile(stream) as z:
        content_types_xml = z.read("[Content_Types].xml").decode("utf-8")
    match = re.search(
        r'<Override PartName="/word/document.xml" ContentType="([^"]+)"', content_types_xml
    )
    assert match is not None, content_types_xml
    return match.group(1)
