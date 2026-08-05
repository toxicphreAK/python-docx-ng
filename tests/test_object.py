# pyright: reportPrivateUsage=false

"""Unit test suite for embedded OLE objects — issue #101."""

from __future__ import annotations

import io

import docx
from docx.object import EmbeddedObject
from docx.opc.constants import CONTENT_TYPE as CT
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import parse_xml
from docx.shared import Inches

from .unitutil.file import test_file

_ICON = test_file("monty-truth.png")
_PAYLOAD = test_file("python-icon.jpeg")


def _reopened(document):
    stream = io.BytesIO()
    document.save(stream)
    return docx.Document(io.BytesIO(stream.getvalue()))


def _document_with_object(**kwargs):
    document = docx.Document()
    run = document.add_paragraph().add_run()
    obj = run.add_embedded_object(_PAYLOAD, icon=_ICON, **kwargs)
    return document, obj


class DescribeAddEmbeddedObject:
    """Unit-test suite for `Run.add_embedded_object()`."""

    def it_embeds_the_file_as_a_part(self):
        document, obj = _document_with_object()

        assert obj.blob is not None
        with open(_PAYLOAD, "rb") as f:
            assert obj.blob == f.read()
        assert obj.content_type == CT.OFC_OLE_OBJECT
        assert obj.filename == "oleObject1.bin"

    def it_writes_the_prog_id_word_launches_on(self):
        _, obj = _document_with_object(prog_id="Excel.Sheet.12")

        assert obj.prog_id == "Excel.Sheet.12"

    def and_defaults_to_the_generic_package_prog_id(self):
        _, obj = _document_with_object()

        assert obj.prog_id == "Package"

    def it_relates_the_icon_image_in_its_own_right(self):
        """An embedded object is at minimum two relationships plus the run XML."""
        document, obj = _document_with_object()

        assert obj.image is not None
        assert obj.image.content_type == "image/png"
        assert len(document.images) == 1

    def it_sizes_the_visual_from_the_icon_by_default(self):
        _, obj = _document_with_object()

        style = obj._object.shape.get("style")
        assert "width:" in style
        assert "height:" in style

    def and_the_size_can_be_given(self):
        _, obj = _document_with_object(width=Inches(0.5), height=Inches(0.5))

        style = obj._object.shape.get("style")
        assert "width:36.00pt" in style
        assert "height:36.00pt" in style

    def it_names_the_shape_the_same_way_in_both_places(self):
        """`o:OLEObject/@ShapeID` must name the `v:shape` or Word cannot find it."""
        _, obj = _document_with_object()

        shape_id = obj._object.shape.get("id")
        assert obj._object.oleObject.ShapeID == shape_id

    def it_shows_the_object_as_an_icon(self):
        _, obj = _document_with_object()

        assert obj.shows_icon is True
        assert obj.is_linked is False

    def it_survives_a_save(self):
        document, _ = _document_with_object(prog_id="Excel.Sheet.12")

        reopened = _reopened(document)

        objects = reopened.embedded_objects
        assert len(objects) == 1
        assert objects[0].prog_id == "Excel.Sheet.12"
        assert objects[0].blob is not None
        assert objects[0].image is not None

    def and_the_payload_lands_under_word_embeddings(self):
        import zipfile

        document, _ = _document_with_object()
        stream = io.BytesIO()
        document.save(stream)

        names = zipfile.ZipFile(io.BytesIO(stream.getvalue())).namelist()
        assert "word/embeddings/oleObject1.bin" in names

    def it_accepts_a_stream(self):
        document = docx.Document()
        run = document.add_paragraph().add_run()
        with open(_PAYLOAD, "rb") as f:
            payload = io.BytesIO(f.read())

        obj = run.add_embedded_object(payload, icon=_ICON)

        assert obj.blob == payload.getvalue()

    def it_numbers_a_second_object_after_the_first(self):
        document = docx.Document()
        run = document.add_paragraph().add_run()

        run.add_embedded_object(_PAYLOAD, icon=_ICON)
        second = run.add_embedded_object(_PAYLOAD, icon=_ICON)

        assert second.filename == "oleObject2.bin"


class DescribeEmbeddedObjectCollections:
    """Unit-test suite for `Run.embedded_objects` and `Document.embedded_objects`."""

    def it_collects_the_objects_of_a_run(self):
        document = docx.Document()
        run = document.add_paragraph().add_run()
        run.add_embedded_object(_PAYLOAD, icon=_ICON)

        assert len(run.embedded_objects) == 1
        assert isinstance(run.embedded_objects[0], EmbeddedObject)

    def and_of_the_document_body(self):
        document = docx.Document()
        for _ in range(2):
            document.add_paragraph().add_run().add_embedded_object(_PAYLOAD, icon=_ICON)

        assert len(document.embedded_objects) == 2

    def and_it_is_empty_for_a_document_with_none(self):
        assert docx.Document().embedded_objects == []


class DescribeReadingAnExistingObject:
    """A document from elsewhere may hold objects this library did not write."""

    def it_reads_a_linked_object_and_reports_no_blob(self):
        """The bytes of a linked object are not in the package at all."""
        document = docx.Document()
        run = document.add_paragraph().add_run()
        run._r.append(
            parse_xml(
                '<w:object %s><o:OLEObject Type="Link" ProgID="Excel.Sheet.12"'
                ' DrawAspect="Icon" r:id="rId99"/></w:object>' % nsdecls("w", "o", "r")
            )
        )

        obj = document.embedded_objects[0]

        assert obj.is_linked is True
        assert obj.prog_id == "Excel.Sheet.12"
        assert obj.blob is None
        assert obj.part_ is None
        assert obj.filename is None

    def it_reports_None_for_an_object_with_no_OLEObject_child(self):
        document = docx.Document()
        run = document.add_paragraph().add_run()
        run._r.append(parse_xml("<w:object %s/>" % nsdecls("w")))

        obj = document.embedded_objects[0]

        assert obj.prog_id is None
        assert obj.blob is None
        assert obj.image is None
        assert obj.shows_icon is False

    def and_for_a_relationship_the_document_does_not_resolve(self):
        document = docx.Document()
        run = document.add_paragraph().add_run()
        run._r.append(
            parse_xml(
                '<w:object %s><o:OLEObject Type="Embed" r:id="rIdMissing"/></w:object>'
                % nsdecls("w", "o", "r")
            )
        )

        assert document.embedded_objects[0].blob is None

    def it_has_a_useful_repr(self):
        _, obj = _document_with_object(prog_id="Excel.Sheet.12")

        assert "Excel.Sheet.12" in repr(obj)


class DescribeOLEElementClasses:
    """VML attribute names are not namespace-qualified and their case is inconsistent."""

    def it_reads_the_unqualified_attributes(self):
        oleObject = parse_xml(
            '<o:OLEObject %s Type="Embed" ProgID="Word.Document.12"'
            ' ShapeID="_x0000_i1025" DrawAspect="Content" ObjectID="_1" r:id="rId4"'
            ' UpdateMode="OnCall"/>' % nsdecls("o", "r")
        )

        assert oleObject.Type == "Embed"
        assert oleObject.ProgID == "Word.Document.12"
        assert oleObject.ShapeID == "_x0000_i1025"
        assert oleObject.DrawAspect == "Content"
        assert oleObject.ObjectID == "_1"
        assert oleObject.rId == "rId4"
        assert oleObject.UpdateMode == "OnCall"

    def it_finds_the_image_relationship_inside_the_shape(self):
        object_elm = parse_xml(
            '<w:object %s><v:shape id="s"><v:imagedata r:id="rId7"/></v:shape>'
            "</w:object>" % nsdecls("w", "v", "r")
        )

        assert object_elm.image_rId == "rId7"

    def and_reports_None_when_there_is_no_shape(self):
        object_elm = parse_xml("<w:object %s/>" % nsdecls("w"))

        assert object_elm.shape is None
        assert object_elm.image_rId is None
        assert object_elm.oleObject is None

    def the_r_id_attribute_is_namespace_qualified_unlike_the_rest(self):
        _, obj = _document_with_object()

        oleObject = obj._object.oleObject
        assert oleObject.get(qn("r:id")) is not None
        assert oleObject.get("ProgID") is not None
