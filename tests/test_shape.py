# pyright: reportPrivateUsage=false

"""Test suite for the docx.shape module."""

from __future__ import annotations

import io
from typing import cast

import pytest

import docx
from docx.document import Document
from docx.enum.shape import WD_INLINE_SHAPE
from docx.oxml.document import CT_Body
from docx.oxml.ns import nsmap
from docx.oxml.shape import CT_Inline
from docx.shape import InlineShape, InlineShapes
from docx.shared import Emu, Inches, Length

from .unitutil.cxml import element, xml
from .unitutil.file import test_file
from .unitutil.mock import FixtureRequest, Mock, instance_mock


class DescribeInlineShapes:
    """Unit-test suite for `docx.shape.InlineShapes` objects."""

    def it_knows_how_many_inline_shapes_it_contains(self, body: CT_Body, document_: Mock):
        inline_shapes = InlineShapes(body, document_)
        assert len(inline_shapes) == 2

    def it_can_iterate_over_its_InlineShape_instances(self, body: CT_Body, document_: Mock):
        inline_shapes = InlineShapes(body, document_)
        assert all(isinstance(s, InlineShape) for s in inline_shapes)
        assert len(list(inline_shapes)) == 2

    def it_provides_indexed_access_to_inline_shapes(self, body: CT_Body, document_: Mock):
        inline_shapes = InlineShapes(body, document_)
        for idx in range(-2, 2):
            assert isinstance(inline_shapes[idx], InlineShape)

    def it_raises_on_indexed_access_out_of_range(self, body: CT_Body, document_: Mock):
        inline_shapes = InlineShapes(body, document_)

        with pytest.raises(IndexError, match=r"inline shape index \[-3\] out of range"):
            inline_shapes[-3]
        with pytest.raises(IndexError, match=r"inline shape index \[2\] out of range"):
            inline_shapes[2]

    def it_knows_the_part_it_belongs_to(self, body: CT_Body, document_: Mock):
        inline_shapes = InlineShapes(body, document_)
        assert inline_shapes.part is document_.part

    # -- fixtures --------------------------------------------------------------------------------

    @pytest.fixture
    def body(self) -> CT_Body:
        return cast(
            CT_Body, element("w:body/w:p/(w:r/w:drawing/wp:inline, w:r/w:drawing/wp:inline)")
        )

    @pytest.fixture
    def document_(self, request: FixtureRequest):
        return instance_mock(request, Document)


class DescribeInlineShape:
    """Unit-test suite for `docx.shape.InlineShape` objects."""

    @pytest.mark.parametrize(
        ("uri", "content_cxml", "expected_value"),
        [
            # -- embedded picture --
            (nsmap["pic"], "/pic:pic/pic:blipFill/a:blip{r:embed=rId1}", WD_INLINE_SHAPE.PICTURE),
            # -- linked picture --
            (
                nsmap["pic"],
                "/pic:pic/pic:blipFill/a:blip{r:link=rId2}",
                WD_INLINE_SHAPE.LINKED_PICTURE,
            ),
            # -- linked and embedded picture (not expected) --
            (
                nsmap["pic"],
                "/pic:pic/pic:blipFill/a:blip{r:embed=rId1,r:link=rId2}",
                WD_INLINE_SHAPE.LINKED_PICTURE,
            ),
            # -- chart --
            (nsmap["c"], "", WD_INLINE_SHAPE.CHART),
            # -- SmartArt --
            (nsmap["dgm"], "", WD_INLINE_SHAPE.SMART_ART),
            # -- something else we don't know about --
            ("foobar", "", WD_INLINE_SHAPE.NOT_IMPLEMENTED),
        ],
    )
    def it_knows_what_type_of_shape_it_is(
        self, uri: str, content_cxml: str, expected_value: WD_INLINE_SHAPE
    ):
        cxml = "wp:inline/a:graphic/a:graphicData{uri=%s}%s" % (uri, content_cxml)
        inline = cast(CT_Inline, element(cxml))
        inline_shape = InlineShape(inline)
        assert inline_shape.type == expected_value

    def it_knows_its_display_dimensions(self):
        inline = cast(CT_Inline, element("wp:inline/wp:extent{cx=333, cy=666}"))
        inline_shape = InlineShape(inline)

        width, height = inline_shape.width, inline_shape.height

        assert isinstance(width, Length)
        assert width == 333
        assert isinstance(height, Length)
        assert height == 666

    def it_can_change_its_display_dimensions(self):
        inline_shape = InlineShape(
            cast(
                CT_Inline,
                element(
                    "wp:inline/(wp:extent{cx=333,cy=666},a:graphic/a:graphicData/pic:pic/"
                    "pic:spPr/a:xfrm/a:ext{cx=333,cy=666})"
                ),
            )
        )

        inline_shape.width = Emu(444)
        inline_shape.height = Emu(888)

        assert inline_shape._inline.xml == xml(
            "wp:inline/(wp:extent{cx=444,cy=888},a:graphic/a:graphicData/pic:pic/pic:spPr/"
            "a:xfrm/a:ext{cx=444,cy=888})"
        )


class DescribeInlineShapeAltText:
    """Unit-test suite for the alternative text of an |InlineShape|."""

    @pytest.mark.parametrize(
        ("docPr_cxml", "expected_description", "expected_title"),
        [
            ("wp:docPr{id=1,name=Picture 1}", None, None),
            ("wp:docPr{id=1,name=Picture 1,descr=A chart}", "A chart", None),
            ("wp:docPr{id=1,name=Picture 1,title=Chart}", None, "Chart"),
            ("wp:docPr{id=1,name=Picture 1,descr=A chart,title=Chart}", "A chart", "Chart"),
        ],
    )
    def it_knows_its_alt_text(
        self, docPr_cxml: str, expected_description: str | None, expected_title: str | None
    ):
        inline_shape = InlineShape(cast(CT_Inline, element("wp:inline/%s" % docPr_cxml)))

        assert inline_shape.description == expected_description
        assert inline_shape.title == expected_title

    def it_can_change_its_alt_text(self):
        inline_shape = InlineShape(
            cast(CT_Inline, element("wp:inline/wp:docPr{id=1,name=Picture 1}"))
        )

        inline_shape.description = "The Python logo"
        inline_shape.title = "Logo"

        assert inline_shape._inline.xml == xml(
            "wp:inline/wp:docPr{id=1,name=Picture 1,descr=The Python logo,title=Logo}"
        )

    def it_can_remove_its_alt_text(self):
        inline_shape = InlineShape(
            cast(
                CT_Inline,
                element("wp:inline/wp:docPr{id=1,name=Picture 1,descr=A chart,title=Chart}"),
            )
        )

        inline_shape.description = None
        inline_shape.title = None

        assert inline_shape._inline.xml == xml("wp:inline/wp:docPr{id=1,name=Picture 1}")

    def it_can_be_given_alt_text_at_insertion_time(self):
        """Requiring a second step to set alt text makes it easy to forget."""
        document = docx.Document()

        document.add_picture(
            test_file("python-icon.jpeg"), description="The Python logo", title="Logo"
        )

        assert document.inline_shapes[0].description == "The Python logo"
        assert document.inline_shapes[0].title == "Logo"

    def it_can_be_given_alt_text_on_a_run(self):
        run = docx.Document().add_paragraph().add_run()

        shape = run.add_picture(test_file("python-icon.jpeg"), description="A snake")

        assert shape.description == "A snake"
        assert shape.title is None


class DescribeSvgPictureInsertion:
    """An SVG picture carries its vector source in an extension of a raster blip."""

    def it_writes_an_svgBlip_extension_alongside_the_blip(self):
        document = docx.Document()

        shape = document.add_picture(test_file("python-logo.svg"))

        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        assert blip.svgBlip is not None
        # -- the extension uri is the fixed GUID a consumer matches on --
        assert blip.xpath("./a:extLst/a:ext/@uri") == ["{96DAC541-7B7A-43D3-8B79-37D633B846F1}"]

    def it_points_the_fallback_blip_at_the_svg_when_no_fallback_is_given(self):
        """Word 2016 and later render this; earlier versions show nothing."""
        document = docx.Document()

        shape = document.add_picture(test_file("python-logo.svg"))

        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        assert blip.embed == blip.svgBlip.embed

    def it_points_the_fallback_blip_at_the_fallback_image_when_given_one(self):
        document = docx.Document()

        shape = document.add_picture(
            test_file("python-logo.svg"), svg_fallback=test_file("monty-truth.png")
        )

        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        assert blip.embed != blip.svgBlip.embed
        # -- the raster part is a second image part, not a second reference to the SVG --
        rels = document.part.rels
        assert rels[blip.embed].target_part.partname.ext == "png"
        assert rels[blip.svgBlip.embed].target_part.partname.ext == "svg"

    def it_sizes_the_picture_from_the_svg_not_the_fallback(self):
        """The SVG states the intended size; the fallback is only a rendering of it."""
        document = docx.Document()

        shape = document.add_picture(
            test_file("python-logo.svg"), svg_fallback=test_file("monty-truth.png")
        )

        assert shape.width == Inches(2)
        assert shape.height == Inches(1)

    def it_writes_no_svgBlip_for_a_raster_picture(self):
        document = docx.Document()

        shape = document.add_picture(test_file("monty-truth.png"))

        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        assert blip.svgBlip is None
        assert blip.embed is not None

    @pytest.mark.parametrize(
        ("filename", "expected_extension"),
        [
            ("python-logo.svg", "svg"),
            ("frame-2x1in.emf", "emf"),
            ("CVS_LOGO.WMF", "WMF"),
        ],
    )
    def it_declares_the_content_type_of_a_vector_part(self, filename: str, expected_extension: str):
        """A part with no content type declared is a document Word refuses to open."""
        import zipfile

        document = docx.Document()
        document.add_picture(test_file(filename))

        stream = io.BytesIO()
        document.save(stream)
        with zipfile.ZipFile(stream) as z:
            content_types = z.read("[Content_Types].xml").decode("utf-8")

        assert 'Extension="%s"' % expected_extension.lower() in content_types

    def it_survives_a_round_trip(self):
        document = docx.Document()
        document.add_picture(
            test_file("python-logo.svg"), svg_fallback=test_file("monty-truth.png")
        )

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        shape = docx.Document(stream).inline_shapes[0]
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        assert blip.svgBlip is not None
        assert shape.width == Inches(2)
