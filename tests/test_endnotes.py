# pyright: reportPrivateUsage=false

"""Unit test suite for the endnote half of the `docx.footnotes` module."""

from __future__ import annotations

import io
from typing import cast

import pytest

import docx
from docx.footnotes import Endnote, Endnotes
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.packuri import PackURI
from docx.oxml.footnotes import CT_Endnotes
from docx.package import Package
from docx.parts.endnotes import EndnotesPart

from .unitutil.cxml import element
from .unitutil.mock import FixtureRequest, Mock, instance_mock

# -- the two endnotes Word keeps for the separator rules --
SEPARATORS = "w:endnote{w:type=separator,w:id=-1},w:endnote{w:type=continuationSeparator,w:id=0}"


def _endnotes(cxml: str, package_: Mock) -> Endnotes:
    endnotes_elm = cast(CT_Endnotes, element(cxml))
    part = EndnotesPart(PackURI("/word/endnotes.xml"), CT.WML_ENDNOTES, endnotes_elm, package_)
    return Endnotes(endnotes_elm, part)


class DescribeEndnotes:
    """Unit-test suite for `docx.footnotes.Endnotes` objects."""

    @pytest.mark.parametrize(
        ("cxml", "count"),
        [
            ("w:endnotes", 0),
            ("w:endnotes/(%s)" % SEPARATORS, 0),
            ("w:endnotes/(%s,w:endnote{w:id=1}/w:p)" % SEPARATORS, 1),
            ("w:endnotes/(%s,w:endnote{w:id=1}/w:p,w:endnote{w:id=2}/w:p)" % SEPARATORS, 2),
        ],
    )
    def it_counts_only_the_endnotes_an_author_wrote(
        self, cxml: str, count: int, package_: Mock
    ):
        """The separators Word keeps at ids -1 and 0 are structural, not content."""
        assert len(_endnotes(cxml, package_)) == count

    def it_can_add_an_endnote(self, package_: Mock):
        endnotes = _endnotes("w:endnotes/(%s)" % SEPARATORS, package_)

        endnote = endnotes.add_endnote()

        # -- ids -1 and 0 are taken by the separators, so authored ids start at 1 --
        assert endnote.endnote_id == 1
        assert len(endnotes) == 1
        # -- the reference mark Word renders as the endnote number is already there --
        assert endnote._note_elm.xpath(".//w:endnoteRef")

    def and_it_applies_the_endnote_styles_rather_than_the_footnote_ones(
        self, package_: Mock
    ):
        endnotes = _endnotes("w:endnotes/(%s)" % SEPARATORS, package_)

        endnote = endnotes.add_endnote()

        assert endnote._note_elm.xpath("./w:p/w:pPr/w:pStyle/@w:val") == ["EndnoteText"]
        assert endnote._note_elm.xpath(".//w:rStyle/@w:val") == ["EndnoteReference"]

    def it_puts_the_first_line_of_text_beside_the_reference_mark(self, package_: Mock):
        endnotes = _endnotes("w:endnotes/(%s)" % SEPARATORS, package_)

        endnote = endnotes.add_endnote("first\nsecond")

        assert endnote.text == "first\nsecond"
        assert len(endnote.paragraphs) == 2

    def it_can_get_an_endnote_by_id(self, package_: Mock):
        endnotes = _endnotes("w:endnotes/(%s,w:endnote{w:id=1}/w:p)" % SEPARATORS, package_)

        endnote = endnotes.get(1)

        assert isinstance(endnote, Endnote)
        assert endnote.endnote_id == 1

    @pytest.mark.parametrize("endnote_id", [-1, 0, 42])
    def but_not_a_separator_or_an_id_that_is_not_there(
        self, endnote_id: int, package_: Mock
    ):
        endnotes = _endnotes("w:endnotes/(%s,w:endnote{w:id=1}/w:p)" % SEPARATORS, package_)

        assert endnotes.get(endnote_id) is None

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def package_(self, request: FixtureRequest):
        return instance_mock(request, Package)


class DescribeEndnotesRoundTrip:
    """The endnotes part is created on demand and survives a save/open cycle."""

    def it_creates_the_endnotes_part_only_when_used(self):
        document = docx.Document()

        assert document.part.has_endnotes_part is False

        document.endnotes

        assert document.part.has_endnotes_part is True

    def it_round_trips_an_endnote_and_its_reference(self):
        document = docx.Document()
        paragraph = document.add_paragraph("body ")
        endnote = document.endnotes.add_endnote("An endnote.")
        paragraph.add_run().add_endnote_reference(endnote)

        stream = io.BytesIO()
        document.save(stream)
        reopened = docx.Document(io.BytesIO(stream.getvalue()))

        assert len(reopened.endnotes) == 1
        assert reopened.endnotes.get(1) is not None
        assert reopened.endnotes.get(1).text == "An endnote."
        assert reopened.paragraphs[0].runs[-1]._r.xpath("./w:endnoteReference/@w:id") == ["1"]

    def it_keeps_footnotes_and_endnotes_apart(self):
        document = docx.Document()
        document.footnotes.add_footnote("foot")
        document.endnotes.add_endnote("end")

        assert [n.text for n in document.footnotes] == ["foot"]
        assert [n.text for n in document.endnotes] == ["end"]

    def it_replaces_text_in_endnotes_when_asked_to_include_footnotes(self):
        """`Document.replace_text()` already documented itself as reaching endnotes."""
        document = docx.Document()
        document.endnotes.add_endnote("needle in an endnote")

        replaced = document.replace_text("needle", "pin", footnotes=True)

        assert replaced == 1
        assert document.endnotes.get(1).text == "pin in an endnote"
