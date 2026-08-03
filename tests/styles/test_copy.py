"""Unit test suite for copying a style between documents."""

from __future__ import annotations

import io

import pytest

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import nsdecls
from docx.oxml.parser import parse_xml
from docx.shared import Pt, RGBColor


def _source_with_callout() -> docx.document.Document:
    """A document defining "Callout", based on "Callout Base" and linked to a char style."""
    document = docx.Document()
    base = document.styles.add_style("Callout Base", WD_STYLE_TYPE.PARAGRAPH)
    base.font.size = Pt(14)
    callout = document.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.base_style = base
    callout.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    linked = document.styles.add_style("Callout Char", WD_STYLE_TYPE.CHARACTER)
    callout._element.link_val = linked.style_id
    return document


def _source_with_list_style(
    fmt: str = "upperRoman", start: int = 1, text: str = "%1."
) -> docx.document.Document:
    """A document defining "Roman List", a paragraph style with a numbering reference."""
    document = docx.Document()
    numbering = document.numbering._element
    numbering.insert(
        0,
        parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="77"><w:lvl w:ilvl="0">'
            f'<w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/></w:lvl></w:abstractNum>'
        ),
    )
    num_id = numbering.add_num(77).numId
    style = document.styles.add_style("Roman List", WD_STYLE_TYPE.PARAGRAPH)
    style._element.numId_val = num_id
    return document


class DescribeCopyStyleFrom:
    """Unit-test suite for `docx.styles.styles.Styles.copy_style_from`."""

    def it_copies_the_style_into_the_document(self):
        source = _source_with_callout()
        destination = docx.Document()
        assert "Callout" not in destination.styles

        copied = destination.styles.copy_style_from(source.styles["Callout"])

        assert copied.name == "Callout"
        assert "Callout" in destination.styles
        assert copied.font.color.rgb == RGBColor(0xC0, 0x00, 0x00)

    def it_makes_the_style_usable_by_name(self):
        """`KeyError: no style with name 'X'` is the problem this exists to solve."""
        source = _source_with_callout()
        destination = docx.Document()

        destination.styles.copy_style_from(source.styles["Callout"])

        paragraph = destination.add_paragraph("text", style="Callout")
        assert paragraph.style.name == "Callout"

    def it_copies_the_based_on_style_too(self):
        """A copied style whose basedOn target is missing renders wrongly."""
        source = _source_with_callout()
        destination = docx.Document()

        copied = destination.styles.copy_style_from(source.styles["Callout"])

        assert "Callout Base" in destination.styles
        assert copied.base_style is not None
        assert copied.base_style.name == "Callout Base"
        assert copied.base_style.font.size == Pt(14)

    def it_copies_the_linked_character_style_too(self):
        source = _source_with_callout()
        destination = docx.Document()

        copied = destination.styles.copy_style_from(source.styles["Callout"])

        link_id = copied._element.link_val
        assert link_id is not None
        linked = destination.styles._element.get_by_id(link_id)
        assert linked is not None, "the w:link target came with it"
        assert linked.name_val == "Callout Char"

    def it_copies_the_next_style_too(self):
        source = docx.Document()
        following = source.styles.add_style("Body After", WD_STYLE_TYPE.PARAGRAPH)
        heading = source.styles.add_style("My Heading", WD_STYLE_TYPE.PARAGRAPH)
        heading.next_paragraph_style = following
        destination = docx.Document()

        copied = destination.styles.copy_style_from(source.styles["My Heading"])

        assert "Body After" in destination.styles
        assert copied.next_paragraph_style.name == "Body After"

    def it_can_be_told_to_skip_the_dependencies(self):
        source = _source_with_callout()
        destination = docx.Document()

        destination.styles.copy_style_from(source.styles["Callout"], include_dependencies=False)

        assert "Callout" in destination.styles
        assert "Callout Base" not in destination.styles

    def it_tolerates_a_dependency_the_source_does_not_define(self):
        """A reference already broken in the source stays broken rather than raising."""
        source = docx.Document()
        style = source.styles.add_style("Orphan", WD_STYLE_TYPE.PARAGRAPH)
        style._element.basedOn_val = "NoSuchStyle"
        destination = docx.Document()

        copied = destination.styles.copy_style_from(source.styles["Orphan"])

        assert copied.name == "Orphan"

    def it_renames_the_style_when_asked(self):
        source = _source_with_callout()
        destination = docx.Document()

        copied = destination.styles.copy_style_from(source.styles["Callout"], name="Sidebar")

        assert copied.name == "Sidebar"
        assert "Sidebar" in destination.styles
        assert "Callout Base" in destination.styles, "dependencies keep their own names"

    def it_gives_the_copy_a_style_id_free_in_the_destination(self):
        source = docx.Document()
        source.styles.add_style("Mine", WD_STYLE_TYPE.PARAGRAPH)
        destination = docx.Document()
        # -- a different style already holding the id the copy would take --
        clash = destination.styles.add_style("Something Else", WD_STYLE_TYPE.PARAGRAPH)
        clash.style_id = "Mine"

        copied = destination.styles.copy_style_from(source.styles["Mine"])

        assert copied.style_id != "Mine"
        assert destination.styles._element.get_by_id("Mine").name_val == "Something Else"

    def it_refuses_to_copy_a_style_into_its_own_document(self):
        document = _source_with_callout()

        with pytest.raises(ValueError, match="already in this document"):
            document.styles.copy_style_from(document.styles["Callout"])

    def it_rejects_an_unknown_collision_policy(self):
        source = _source_with_callout()

        with pytest.raises(ValueError, match="on_collision must be one of"):
            docx.Document().styles.copy_style_from(source.styles["Callout"], on_collision="clobber")

    def it_survives_a_round_trip(self):
        source = _source_with_callout()
        destination = docx.Document()
        destination.styles.copy_style_from(source.styles["Callout"])
        destination.add_paragraph("text", style="Callout")

        stream = io.BytesIO()
        destination.save(stream)
        stream.seek(0)
        reopened = docx.Document(stream)

        assert reopened.paragraphs[0].style.name == "Callout"
        assert reopened.styles["Callout"].base_style.name == "Callout Base"


class DescribeCopyStyleCollisions:
    """Unit-test suite for the `on_collision` policies."""

    @pytest.fixture
    def documents(self):
        source = docx.Document()
        style = source.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        style.font.size = Pt(20)

        destination = docx.Document()
        existing = destination.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        existing.font.size = Pt(8)
        return source, destination

    def it_skips_by_default_and_returns_the_existing_style(self, documents):
        source, destination = documents

        copied = destination.styles.copy_style_from(source.styles["Callout"])

        assert copied.font.size == Pt(8), "the existing definition is untouched"
        assert len([s for s in destination.styles if s.name == "Callout"]) == 1

    def it_can_overwrite_the_existing_definition(self, documents):
        source, destination = documents

        copied = destination.styles.copy_style_from(
            source.styles["Callout"], on_collision="overwrite"
        )

        assert copied.font.size == Pt(20)
        assert len([s for s in destination.styles if s.name == "Callout"]) == 1

    def it_can_copy_under_a_free_name(self, documents):
        source, destination = documents

        copied = destination.styles.copy_style_from(source.styles["Callout"], on_collision="rename")

        assert copied.name == "Callout 2"
        assert copied.font.size == Pt(20)
        assert destination.styles["Callout"].font.size == Pt(8)

    def it_can_raise_on_a_collision(self, documents):
        source, destination = documents

        with pytest.raises(ValueError, match="already contains style 'Callout'"):
            destination.styles.copy_style_from(source.styles["Callout"], on_collision="raise")

    def it_reuses_an_existing_dependency_whatever_the_policy_says(self):
        """Renaming or overwriting "Normal" because a copied style is based on it would
        be a surprising thing to do to the destination document."""
        source = _source_with_callout()
        destination = docx.Document()
        destination.styles.add_style("Callout Base", WD_STYLE_TYPE.PARAGRAPH)

        destination.styles.copy_style_from(source.styles["Callout"], on_collision="rename")

        base_styles = [s for s in destination.styles if s.name.startswith("Callout Base")]
        assert len(base_styles) == 1, "the dependency was reused, not duplicated"


class DescribeCopyStyleNumbering:
    """Unit-test suite for carrying a list style's numbering across."""

    def it_copies_the_numbering_definition_and_repoints_the_style(self):
        source = _source_with_list_style(fmt="upperRoman", start=3, text="[%1]")
        destination = docx.Document()
        before_abstract = {a.abstractNumId for a in destination.numbering._element.abstractNum_lst}
        before_num = {n.numId for n in destination.numbering._element.num_lst}

        copied = destination.styles.copy_style_from(source.styles["Roman List"])

        after_abstract = {a.abstractNumId for a in destination.numbering._element.abstractNum_lst}
        after_num = {n.numId for n in destination.numbering._element.num_lst}
        assert len(after_abstract - before_abstract) == 1, "a new abstract definition"
        assert copied._element.numId_val in (after_num - before_num)

    def it_keeps_the_numbering_part_in_schema_order(self):
        """`CT_Numbering` is an xsd:sequence: every abstractNum precedes every num, and
        Word refuses to open a document that gets it wrong."""
        source = _source_with_list_style()
        destination = docx.Document()

        destination.styles.copy_style_from(source.styles["Roman List"])

        local_names = [c.tag.split("}")[1] for c in destination.numbering._element]
        first_num = local_names.index("num")
        assert "abstractNum" not in local_names[first_num:]

    def it_numbers_correctly_in_the_destination(self):
        source = _source_with_list_style(fmt="upperRoman", start=3, text="[%1]")
        destination = docx.Document()

        destination.styles.copy_style_from(source.styles["Roman List"])
        for text in ("a", "b", "c"):
            destination.add_paragraph(text, style="Roman List")

        assert [n for _, n in destination.list_numbers] == ["[III]", "[IV]", "[V]"]

    def it_can_be_told_to_leave_numbering_alone(self):
        source = _source_with_list_style()
        destination = docx.Document()
        before = len(destination.numbering._element.abstractNum_lst)

        destination.styles.copy_style_from(source.styles["Roman List"], include_numbering=False)

        assert len(destination.numbering._element.abstractNum_lst) == before

    def it_copies_one_abstract_definition_once_for_several_styles(self):
        source = _source_with_list_style()
        shared_num_id = source.styles["Roman List"]._element.numId_val
        second = source.styles.add_style("Roman List B", WD_STYLE_TYPE.PARAGRAPH)
        second._element.numId_val = shared_num_id
        second.base_style = source.styles["Roman List"]
        destination = docx.Document()
        before = len(destination.numbering._element.abstractNum_lst)

        destination.styles.copy_style_from(source.styles["Roman List B"])

        added = len(destination.numbering._element.abstractNum_lst) - before
        assert added == 1, "both styles share the one abstract definition"

    def it_tolerates_a_style_whose_numbering_reference_is_already_dangling(self):
        source = docx.Document()
        style = source.styles.add_style("Broken List", WD_STYLE_TYPE.PARAGRAPH)
        style._element.numId_val = 9999
        destination = docx.Document()

        copied = destination.styles.copy_style_from(source.styles["Broken List"])

        assert copied.name == "Broken List"

    def it_survives_a_round_trip(self):
        source = _source_with_list_style(fmt="lowerLetter", text="%1)")
        destination = docx.Document()
        destination.styles.copy_style_from(source.styles["Roman List"])
        for text in ("a", "b"):
            destination.add_paragraph(text, style="Roman List")

        stream = io.BytesIO()
        destination.save(stream)
        stream.seek(0)

        assert [n for _, n in docx.Document(stream).list_numbers] == ["a)", "b)"]
