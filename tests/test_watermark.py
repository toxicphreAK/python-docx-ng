"""Unit test suite for the docx.watermark module."""

from __future__ import annotations

import io

import pytest

import docx
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

_IMAGE = "tests/test_files/python-icon.png"


def _shape_style(watermark) -> dict[str, str]:
    """The VML style string of `watermark` as a dict of its declarations."""
    style = watermark._shape.get("style") or ""
    return dict(
        part.split(":", 1)
        for part in style.split(";")
        if ":" in part  # pyright: ignore
    )


class DescribeAddTextWatermark:
    """Unit-test suite for text watermarks."""

    def it_adds_the_watermark_to_all_three_header_types(self):
        """A watermark that vanishes on page 1 reads as a bug rather than a setting."""
        document = docx.Document()

        watermarks = document.add_text_watermark("DRAFT")

        assert len(watermarks) == 3
        assert [w.text for w in watermarks] == ["DRAFT"] * 3
        section = document.sections[0]
        for header in (section.header, section.first_page_header, section.even_page_header):
            assert len(header.part.element.xpath(".//w:pict/v:shape")) == 1

    def it_writes_the_shape_type_the_shape_refers_to(self):
        """The WordArt shape does not render without its shape-type definition."""
        document = docx.Document()

        document.add_text_watermark("DRAFT")

        hdr = document.sections[0].header.part.element
        assert hdr.xpath('.//v:shapetype[@id="_x0000_t136"]')
        assert hdr.xpath('.//v:shape[@type="#_x0000_t136"]')

    def it_names_the_shape_the_way_word_does(self):
        """Word finds an existing watermark by this name, including to remove it."""
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark("DRAFT")

        assert watermark._shape.get("id") == "PowerPlusWaterMarkObject"

    def it_places_the_watermark_behind_the_text(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark("DRAFT")

        assert int(_shape_style(watermark)["z-index"]) < 0

    def it_centres_the_watermark_on_the_page(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark("DRAFT")

        style = _shape_style(watermark)
        assert style["mso-position-horizontal"] == "center"
        assert style["mso-position-vertical"] == "center"

    def it_defaults_to_words_own_size_and_diagonal_angle(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark("DRAFT")

        style = _shape_style(watermark)
        assert style["rotation"] == "315"
        assert style["width"] == "468pt"
        assert style["height"] == "234pt"

    def it_accepts_a_horizontal_angle_and_a_custom_size(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark(
            "DRAFT", angle=0, width=Pt(300), height=Pt(100)
        )

        style = _shape_style(watermark)
        assert style["rotation"] == "0"
        assert (style["width"], style["height"]) == ("300pt", "100pt")

    def it_accepts_a_colour_with_or_without_a_leading_hash(self):
        document = docx.Document()

        (a, *_) = document.add_text_watermark("A", color="FF0000")
        (b, *_) = document.sections[0].add_text_watermark("B", color="#00FF00")

        assert a._shape.get("fillcolor") == "#FF0000"
        assert b._shape.get("fillcolor") == "#00FF00"

    def it_stretches_the_text_to_the_shape_unless_given_a_font_size(self):
        document = docx.Document()

        (stretched, *_) = document.add_text_watermark("A")
        (sized, *_) = document.sections[0].add_text_watermark("B", font_size=Pt(36))

        stretched_path = stretched._shape.find(qn("v:textpath"))
        sized_path = sized._shape.find(qn("v:textpath"))
        assert "font-size:1pt" in stretched_path.get("style")
        assert stretched_path.get("fitshape") is None, "the shape type's fitshape applies"
        assert "font-size:36pt" in sized_path.get("style")
        assert sized_path.get("fitshape") == "f"

    def it_carries_the_font_and_emphasis_through(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark("DRAFT", font="Arial", bold=True, italic=True)

        # -- the `&quot;` written into the style is decoded on parse, so the attribute
        # -- value holds real quotes; lxml re-escapes them on serialization --
        style = watermark._shape.find(qn("v:textpath")).get("style")
        assert 'font-family:"Arial"' in style
        assert "font-weight:bold" in style
        assert "font-style:italic" in style

    def it_can_set_true_vml_opacity(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark("DRAFT", opacity=0.5)

        fill = watermark._shape.find(qn("v:fill"))
        assert fill is not None
        assert fill.get("opacity") == "0.5"

    def it_escapes_text_that_would_otherwise_break_the_attribute(self):
        document = docx.Document()

        (watermark, *_) = document.add_text_watermark('a"b & <c>')

        assert watermark.text == 'a"b & <c>'

    def it_contributes_no_text_to_the_header(self):
        document = docx.Document()

        document.add_text_watermark("DRAFT")

        assert document.sections[0].header.paragraphs[0].text == ""


class DescribeAddImageWatermark:
    """Unit-test suite for image watermarks."""

    def it_adds_an_image_watermark_to_all_three_header_types(self):
        document = docx.Document()

        watermarks = document.add_image_watermark(_IMAGE)

        assert len(watermarks) == 3
        assert all(w.is_image for w in watermarks)
        assert all(w.text is None for w in watermarks)

    def it_relates_the_image_to_each_header_part(self):
        """A relationship belongs to the part that refers to it."""
        document = docx.Document()

        document.add_image_watermark(_IMAGE)

        section = document.sections[0]
        for header in (section.header, section.first_page_header, section.even_page_header):
            (imagedata,) = header.part.element.xpath(".//v:imagedata")
            rId = imagedata.get(qn("r:id"))
            assert rId in header.part.rels

    def it_applies_words_washout_correction_by_default(self):
        document = docx.Document()

        (watermark, *_) = document.add_image_watermark(_IMAGE)

        imagedata = watermark._shape.find(qn("v:imagedata"))
        assert imagedata.get("gain") == "19661f"
        assert imagedata.get("blacklevel") == "22938f"

    def it_can_leave_the_image_at_full_strength(self):
        document = docx.Document()

        (watermark, *_) = document.add_image_watermark(_IMAGE, washout=False)

        imagedata = watermark._shape.find(qn("v:imagedata"))
        assert imagedata.get("gain") is None

    def it_scales_the_image(self):
        document = docx.Document()

        (unscaled, *_) = document.add_image_watermark(_IMAGE, width=Cm(4))
        document.remove_watermark()
        (scaled, *_) = document.add_image_watermark(_IMAGE, width=Cm(4), scale=0.5)

        assert _shape_style(unscaled)["width"] == "113.39pt"
        assert _shape_style(scaled)["width"] == "56.69pt"

    def it_writes_the_picture_shape_type(self):
        document = docx.Document()

        document.add_image_watermark(_IMAGE)

        hdr = document.sections[0].header.part.element
        assert hdr.xpath('.//v:shapetype[@id="_x0000_t75"]')


class DescribeWatermarkScope:
    """Unit-test suite for which headers a watermark reaches."""

    def it_writes_once_to_a_header_shared_between_sections(self):
        document = docx.Document()
        document.add_paragraph("one")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("two")

        watermarks = document.add_text_watermark("DRAFT")

        assert len(document.sections) == 2
        assert len(watermarks) == 3, "the second section inherits all three headers"

    def it_writes_to_a_section_that_has_its_own_header(self):
        document = docx.Document()
        document.add_paragraph("one")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("two")
        document.sections[1].header.is_linked_to_previous = False

        watermarks = document.add_text_watermark("DRAFT")

        assert len(watermarks) == 4

    def it_can_watermark_a_single_section(self):
        document = docx.Document()
        document.add_paragraph("one")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.sections[1].header.is_linked_to_previous = False

        watermarks = document.sections[1].add_text_watermark("DRAFT")

        assert len(watermarks) == 3
        assert len(document.sections[1].watermarks) == 3


class DescribeRemoveWatermark:
    """Unit-test suite for watermark removal."""

    def it_removes_every_watermark_and_reports_how_many(self):
        document = docx.Document()
        document.add_text_watermark("DRAFT")

        removed = document.remove_watermark()

        assert removed == 3
        assert document.watermarks == []

    def it_leaves_no_empty_run_behind(self):
        document = docx.Document()
        document.add_text_watermark("DRAFT")

        document.remove_watermark()

        hdr = document.sections[0].header.part.element
        assert hdr.xpath(".//w:r") == []

    def it_removes_nothing_from_a_document_with_no_watermark(self):
        document = docx.Document()

        assert document.remove_watermark() == 0

    def it_can_remove_one_watermark_on_its_own(self):
        document = docx.Document()
        watermarks = document.add_text_watermark("DRAFT")

        watermarks[0].remove()

        assert len(document.watermarks) == 2

    def it_can_remove_a_watermark_from_one_section(self):
        document = docx.Document()
        document.add_paragraph("one")
        document.add_section(WD_SECTION.NEW_PAGE)
        document.sections[1].header.is_linked_to_previous = False
        document.add_text_watermark("DRAFT")

        removed = document.sections[1].remove_watermark()

        assert removed == 3, "its own header plus the two it inherits"


class DescribeWatermarkRoundTrip:
    """Unit-test suite for reading a watermark back out of a saved document."""

    @pytest.mark.parametrize("text", ["DRAFT", "CONFIDENTIAL", 'a"b & <c>'])
    def it_survives_a_round_trip(self, text: str):
        document = docx.Document()
        document.add_text_watermark(text)

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        reopened = docx.Document(stream)

        assert [w.text for w in reopened.watermarks] == [text] * 3

    def it_survives_a_round_trip_as_an_image(self):
        document = docx.Document()
        document.add_image_watermark(_IMAGE)

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        reopened = docx.Document(stream)

        assert len(reopened.watermarks) == 3
        assert all(w.is_image for w in reopened.watermarks)

    def it_can_be_removed_after_a_round_trip(self):
        document = docx.Document()
        document.add_text_watermark("DRAFT")
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        reopened = docx.Document(stream)
        assert reopened.remove_watermark() == 3

        stream2 = io.BytesIO()
        reopened.save(stream2)
        stream2.seek(0)
        assert docx.Document(stream2).watermarks == []
