"""Integration tests for byte-reproducible document output.

Saving the same content twice must produce identical bytes, so that build pipelines can
cache on, sign, or diff generated documents.
"""

from __future__ import annotations

import io
import zipfile

import docx


def _build_document() -> bytes:
    """Return the serialized bytes of a document exercising several part types."""
    document = docx.Document()
    document.add_heading("Reproducible", level=1)
    document.add_paragraph("It was a dark and stormy night.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(1, 1).text = "b"
    document.add_page_break()

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class DescribeReproducibleOutput:
    """Acceptance suite for deterministic package serialization."""

    def it_writes_identical_bytes_for_identical_content(self):
        assert _build_document() == _build_document()

    def it_writes_parts_in_partname_order(self):
        with zipfile.ZipFile(io.BytesIO(_build_document())) as z:
            membernames = z.namelist()

        # -- the content-types stream is written first by definition, the rest follow
        # -- in sorted partname order --
        assert membernames[0] == "[Content_Types].xml"
        parts = [n for n in membernames if not n.endswith(".rels")]
        assert parts == sorted(parts)

    def it_stamps_every_member_with_a_fixed_timestamp(self):
        with zipfile.ZipFile(io.BytesIO(_build_document())) as z:
            date_times = {i.date_time for i in z.infolist()}

        assert date_times == {(1980, 1, 1, 0, 0, 0)}

    def it_compresses_every_member(self):
        with zipfile.ZipFile(io.BytesIO(_build_document())) as z:
            compress_types = {i.compress_type for i in z.infolist()}

        assert compress_types == {zipfile.ZIP_DEFLATED}

    def it_round_trips_a_saved_document_to_identical_bytes(self):
        """Loading a generated document and re-saving it must not change the bytes."""
        original = _build_document()

        stream = io.BytesIO()
        docx.Document(io.BytesIO(original)).save(stream)

        assert stream.getvalue() == original
