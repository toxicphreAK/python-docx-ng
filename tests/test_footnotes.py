# pyright: reportPrivateUsage=false

"""Unit test suite for the `docx.footnotes` module."""

from __future__ import annotations

from typing import cast

import pytest

from docx.footnotes import Footnote, Footnotes
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.packuri import PackURI
from docx.oxml.footnotes import CT_Footnotes, CT_FtnEdn
from docx.oxml.ns import qn
from docx.oxml.text.run import CT_R
from docx.package import Package
from docx.parts.document import DocumentPart
from docx.parts.footnotes import FootnotesPart
from docx.text.run import Run

from .unitutil.cxml import element
from .unitutil.mock import FixtureRequest, Mock, instance_mock

# -- the two footnotes Word keeps for the separator rules, as they appear in a real
# -- footnotes part --
SEPARATORS = "w:footnote{w:type=separator,w:id=-1},w:footnote{w:type=continuationSeparator,w:id=0}"


def _footnotes(cxml: str, package_: Mock) -> Footnotes:
    footnotes_elm = cast(CT_Footnotes, element(cxml))
    part = FootnotesPart(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, footnotes_elm, package_)
    return Footnotes(footnotes_elm, part)


class DescribeFootnotes:
    """Unit-test suite for `docx.footnotes.Footnotes` objects."""

    @pytest.mark.parametrize(
        ("cxml", "count"),
        [
            ("w:footnotes", 0),
            ("w:footnotes/(%s)" % SEPARATORS, 0),
            ("w:footnotes/(%s,w:footnote{w:id=1})" % SEPARATORS, 1),
            ("w:footnotes/(%s,w:footnote{w:id=1},w:footnote{w:id=2})" % SEPARATORS, 2),
            # -- an explicit `w:type` of "normal" is an authored footnote --
            ("w:footnotes/w:footnote{w:id=1,w:type=normal}", 1),
        ],
    )
    def it_counts_only_the_footnotes_an_author_wrote(self, cxml: str, count: int, package_: Mock):
        assert len(_footnotes(cxml, package_)) == count

    def it_is_iterable_over_the_footnotes_it_contains(self, package_: Mock):
        footnotes = _footnotes(
            "w:footnotes/(%s,w:footnote{w:id=1},w:footnote{w:id=2})" % SEPARATORS, package_
        )

        items = list(footnotes)

        assert [f.footnote_id for f in items] == [1, 2]
        assert all(isinstance(f, Footnote) for f in items)

    @pytest.mark.parametrize(
        ("footnote_id", "expected_id"), [(1, 1), (2, 2), (-1, None), (0, None), (99, None)]
    )
    def it_can_get_a_footnote_by_id(
        self, footnote_id: int, expected_id: int | None, package_: Mock
    ):
        footnotes = _footnotes(
            "w:footnotes/(%s,w:footnote{w:id=1},w:footnote{w:id=2})" % SEPARATORS, package_
        )

        footnote = footnotes.get(footnote_id)

        assert (footnote.footnote_id if footnote is not None else None) == expected_id

    def it_can_add_a_footnote(self, package_: Mock):
        footnotes = _footnotes("w:footnotes/(%s)" % SEPARATORS, package_)

        footnote = footnotes.add_footnote()

        # -- ids -1 and 0 are taken by the separators, so authored ids start at 1 --
        assert footnote.footnote_id == 1
        assert len(footnotes) == 1
        # -- the reference mark Word renders as the footnote number is already there --
        assert footnote._footnote_elm.xpath(".//w:footnoteRef")

    def and_it_numbers_each_new_footnote_after_the_last(self, package_: Mock):
        footnotes = _footnotes("w:footnotes/(%s,w:footnote{w:id=7})" % SEPARATORS, package_)

        assert [footnotes.add_footnote().footnote_id for _ in range(2)] == [8, 9]

    def and_it_can_add_a_footnote_with_text(self, package_: Mock):
        footnotes = _footnotes("w:footnotes/(%s)" % SEPARATORS, package_)

        footnote = footnotes.add_footnote("See Smith (2019).")

        assert footnote.text == "See Smith (2019)."
        # -- the text joins the paragraph holding the reference mark --
        assert len(footnote.paragraphs) == 1

    def and_a_newline_in_the_text_starts_a_new_paragraph(self, package_: Mock):
        footnotes = _footnotes("w:footnotes/(%s)" % SEPARATORS, package_)

        footnote = footnotes.add_footnote("First line.\nSecond line.")

        assert footnote.text == "First line.\nSecond line."
        assert len(footnote.paragraphs) == 2


class DescribeFootnote:
    """Unit-test suite for `docx.footnotes.Footnote` objects."""

    def it_knows_its_id(self, package_: Mock):
        footnote = list(_footnotes("w:footnotes/w:footnote{w:id=4}", package_))[0]

        assert footnote.footnote_id == 4

    def it_knows_its_text(self, package_: Mock):
        cxml = 'w:footnotes/w:footnote{w:id=1}/(w:p/w:r/w:t"one",w:p/w:r/w:t"two")'
        footnote = list(_footnotes(cxml, package_))[0]

        assert footnote.text == "one\ntwo"

    def it_applies_the_footnote_text_style_to_a_paragraph_it_adds(self, package_: Mock):
        footnote = list(_footnotes("w:footnotes/w:footnote{w:id=1}", package_))[0]

        paragraph = footnote.add_paragraph("A note.")

        assert paragraph._p.style == "FootnoteText"

    def but_an_explicit_style_wins(self, package_: Mock):
        package_.main_document_part.get_style_id.return_value = "Quote"
        footnote = list(_footnotes("w:footnotes/w:footnote{w:id=1}", package_))[0]

        paragraph = footnote.add_paragraph("A note.", "Quote")

        assert paragraph._p.style == "Quote"


class DescribeCT_FtnEdn:
    """Unit-test suite for `docx.oxml.footnotes.CT_FtnEdn`."""

    @pytest.mark.parametrize(
        ("cxml", "expected_value"),
        [
            ("w:footnote{w:id=1}", False),
            ("w:footnote{w:id=1,w:type=normal}", False),
            ("w:footnote{w:id=-1,w:type=separator}", True),
            ("w:footnote{w:id=0,w:type=continuationSeparator}", True),
            ("w:footnote{w:id=2,w:type=continuationNotice}", True),
        ],
    )
    def it_knows_whether_it_is_one_of_words_structural_footnotes(
        self, cxml: str, expected_value: bool
    ):
        footnote = cast(CT_FtnEdn, element(cxml))

        assert footnote.is_structural is expected_value


class DescribeRunAddFootnoteReference:
    """Unit-test suite for `docx.text.run.Run.add_footnote_reference`."""

    def it_adds_a_reference_to_the_footnote(self, package_: Mock, parent_: Mock):
        footnote = list(_footnotes("w:footnotes/w:footnote{w:id=3}", package_))[0]
        run = Run(cast(CT_R, element("w:r")), parent_)

        run.add_footnote_reference(footnote)

        ref = run._r.find(qn("w:footnoteReference"))
        assert ref is not None
        assert ref.get(qn("w:id")) == "3"
        # -- Word relies on this style to raise the mark to a superscript --
        assert run._r.style == "FootnoteReference"

    def but_it_leaves_a_run_that_already_has_a_style_alone(self, package_: Mock, parent_: Mock):
        footnote = list(_footnotes("w:footnotes/w:footnote{w:id=3}", package_))[0]
        run = Run(cast(CT_R, element("w:r/w:rPr/w:rStyle{w:val=Emphasis}")), parent_)

        run.add_footnote_reference(footnote)

        assert run._r.style == "Emphasis"

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def parent_(self, request: FixtureRequest):
        return instance_mock(request, DocumentPart)


@pytest.fixture
def package_(request: FixtureRequest):
    return instance_mock(request, Package)
