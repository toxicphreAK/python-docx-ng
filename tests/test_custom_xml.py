# pyright: reportPrivateUsage=false

"""Unit test suite for the custom XML data store parts."""

from __future__ import annotations

import io

import docx
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.parts.custom_xml import CustomXmlPart, CustomXmlPropertiesPart
from docx.oxml.customxml import CT_DatastoreItem

from .unitutil.cxml import element

_INVOICE = "<invoice><total>42.00</total></invoice>"


class DescribeCustomXmlParts:
    """The bundled template carries a bibliography store, so a default document has one."""

    def it_exposes_the_data_store_items_of_a_document(self):
        document = docx.Document()

        parts = document.custom_xml_parts

        assert len(parts) == 1
        assert isinstance(parts[0], CustomXmlPart)
        assert str(parts[0].partname) == "/customXml/item1.xml"

    def it_reads_the_item_id_and_schema_refs_from_the_props_sidecar(self):
        part = docx.Document().custom_xml_parts[0]

        assert part.item_id == "{EF278816-EC6F-A645-907D-7F25AECB1D4A}"
        assert part.schema_refs == (
            "http://schemas.openxmlformats.org/officeDocument/2006/bibliography",
        )

    def it_exposes_the_item_xml(self):
        part = docx.Document().custom_xml_parts[0]

        assert part.xml.startswith("<b:Sources")
        assert part.element is not None

    def it_types_the_props_part_as_a_CustomXmlPropertiesPart(self):
        document = docx.Document()

        props = [
            p
            for p in document.part.package.iter_parts()
            if isinstance(p, CustomXmlPropertiesPart)
        ]

        assert len(props) == 1


class DescribeAddCustomXmlPart:
    """Unit-test suite for `Document.add_custom_xml_part()`."""

    def it_adds_an_item_and_its_props_sidecar(self):
        document = docx.Document()

        part = document.add_custom_xml_part(_INVOICE, ("urn:example:invoice",))

        assert str(part.partname) == "/customXml/item2.xml"
        assert part.schema_refs == ("urn:example:invoice",)
        # -- the sidecar is numbered after its item, as Word pairs them --
        props = part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps"
        )
        assert str(props.partname) == "/customXml/itemProps2.xml"
        assert props.content_type == CT.OFC_CUSTOM_XML_PROPERTIES

    def it_generates_a_guid_item_id(self):
        document = docx.Document()

        part = document.add_custom_xml_part(_INVOICE)

        assert part.item_id is not None
        assert part.item_id.startswith("{")
        assert part.item_id.endswith("}")
        assert len(part.item_id) == 38
        assert part.item_id != document.custom_xml_parts[0].item_id

    def it_accepts_bytes_as_well_as_a_string(self):
        document = docx.Document()

        part = document.add_custom_xml_part(_INVOICE.encode("utf-8"))

        assert "42.00" in part.xml

    def it_round_trips_through_a_save(self):
        document = docx.Document()
        document.add_custom_xml_part(_INVOICE, ("urn:example:invoice",))
        stream = io.BytesIO()
        document.save(stream)

        reopened = docx.Document(io.BytesIO(stream.getvalue()))

        parts = reopened.custom_xml_parts
        assert len(parts) == 2
        added = parts[-1]
        assert "42.00" in added.xml
        assert added.schema_refs == ("urn:example:invoice",)

    def it_leaves_an_existing_store_undisturbed_on_open_and_save(self):
        """A data store is an opaque payload; merely rewriting must not touch it."""
        document = docx.Document()
        before = document.custom_xml_parts[0].xml

        stream = io.BytesIO()
        document.save(stream)
        reopened = docx.Document(io.BytesIO(stream.getvalue()))

        assert reopened.custom_xml_parts[0].xml == before
        assert reopened.custom_xml_parts[0].item_id == (
            "{EF278816-EC6F-A645-907D-7F25AECB1D4A}"
        )


class DescribeCT_DatastoreItem:
    """Unit-test suite for the `ds:datastoreItem` element class."""

    def it_can_be_created_with_schema_refs(self):
        datastoreItem = CT_DatastoreItem.new("{ABC}", ("urn:a", "urn:b"))

        assert datastoreItem.itemID == "{ABC}"
        assert datastoreItem.schema_ref_uris == ("urn:a", "urn:b")

    def and_without_them(self):
        datastoreItem = CT_DatastoreItem.new("{ABC}")

        assert datastoreItem.schemaRefs is None
        assert datastoreItem.schema_ref_uris == ()

    def it_reads_schema_refs_from_an_existing_element(self):
        datastoreItem = element(
            "ds:datastoreItem{ds:itemID=X}/ds:schemaRefs/ds:schemaRef{ds:uri=urn:a}"
        )

        assert datastoreItem.schema_ref_uris == ("urn:a",)
