"""Test suite for the docx.opc.extendedprops module."""

from __future__ import annotations

import io

import pytest

import docx
from docx.opc.extendedprops import ExtendedProperties
from docx.oxml.extendedprops import CT_ExtendedProperties


def _props(*children: str) -> ExtendedProperties:
    """Return |ExtendedProperties| on a `ep:Properties` element having `children`."""
    element = CT_ExtendedProperties.new()
    for child_xml in children:
        name, _, text = child_xml.partition("=")
        child = getattr(element, "get_or_add_%s" % name)()
        child.text = text
    return ExtendedProperties(element)


class DescribeExtendedProperties:
    """Unit-test suite for `docx.opc.extendedprops.ExtendedProperties`."""

    @pytest.mark.parametrize(
        ("prop_name", "element_name", "value"),
        [
            ("template", "Template", "Normal.dotm"),
            ("manager", "Manager", "A. Manager"),
            ("company", "Company", "Acme Corp"),
            ("application", "Application", "python-docx-ng"),
            ("app_version", "AppVersion", "16.0000"),
            ("presentation_format", "PresentationFormat", "A4 Paper"),
            ("hyperlink_base", "HyperlinkBase", "https://example.com/"),
        ],
    )
    def it_reads_its_string_properties(self, prop_name: str, element_name: str, value: str):
        props = _props("%s=%s" % (element_name, value))
        assert getattr(props, prop_name) == value

    @pytest.mark.parametrize(
        ("prop_name", "value"),
        [
            ("template", "Other.dotx"),
            ("manager", "B. Manager"),
            ("company", "Other Corp"),
            ("application", "Some App"),
            ("app_version", "12.0000"),
            ("presentation_format", "Letter"),
            ("hyperlink_base", "https://other.example/"),
        ],
    )
    def it_can_change_its_string_properties(self, prop_name: str, value: str):
        props = _props()

        setattr(props, prop_name, value)

        assert getattr(props, prop_name) == value

    @pytest.mark.parametrize(
        ("prop_name", "element_name"),
        [
            ("pages", "Pages"),
            ("words", "Words"),
            ("characters", "Characters"),
            ("characters_with_spaces", "CharactersWithSpaces"),
            ("lines", "Lines"),
            ("paragraphs", "Paragraphs"),
            ("total_time", "TotalTime"),
            ("doc_security", "DocSecurity"),
        ],
    )
    def it_reads_and_writes_its_int_properties(self, prop_name: str, element_name: str):
        props = _props("%s=42" % element_name)
        assert getattr(props, prop_name) == 42

        setattr(props, prop_name, 7)

        assert getattr(props, prop_name) == 7

    @pytest.mark.parametrize(
        ("prop_name", "element_name"),
        [
            ("scale_crop", "ScaleCrop"),
            ("links_up_to_date", "LinksUpToDate"),
            ("shared_doc", "SharedDoc"),
            ("hyperlinks_changed", "HyperlinksChanged"),
        ],
    )
    def it_reads_and_writes_its_bool_properties(self, prop_name: str, element_name: str):
        assert getattr(_props("%s=true" % element_name), prop_name) is True
        assert getattr(_props("%s=false" % element_name), prop_name) is False

        props = _props()
        setattr(props, prop_name, True)

        assert getattr(props, prop_name) is True
        # -- serialized as the XSD boolean "true", not Python's "True" --
        assert "<ep:%s>true</ep:%s>" % (element_name, element_name) in props._element.xml

    def it_reports_None_for_an_absent_property(self):
        props = _props()

        assert props.company is None
        assert props.words is None
        assert props.scale_crop is None

    def it_distinguishes_an_absent_property_from_an_empty_one(self):
        """Word writes empty elements like `<HyperlinkBase/>`; that is not "unset"."""
        assert _props().hyperlink_base is None
        assert _props("HyperlinkBase=").hyperlink_base == ""

    def it_removes_the_element_when_assigned_None(self):
        props = _props("Company=Acme Corp")

        props.company = None

        assert props.company is None
        assert "Company" not in props._element.xml

    def it_tolerates_a_non_numeric_value_in_an_int_property(self):
        """These values come from other applications and must not break loading."""
        assert _props("Words=lots").words is None


class DescribeExtendedPropertiesIntegration:
    """Integration-test suite for extended properties on a real package."""

    def it_provides_access_to_the_extended_properties_of_a_document(self):
        document = docx.Document()

        # -- values carried by the default template --
        assert document.extended_properties.template == "Normal.dotm"
        assert document.extended_properties.application == "Microsoft Macintosh Word"

    def it_round_trips_extended_properties_through_a_save(self):
        document = docx.Document()
        document.extended_properties.company = "Acme Corp"
        document.extended_properties.words = 1234
        document.extended_properties.total_time = 42
        document.extended_properties.scale_crop = True

        stream = io.BytesIO()
        document.save(stream)
        reloaded = docx.Document(stream).extended_properties

        assert reloaded.company == "Acme Corp"
        assert reloaded.words == 1234
        assert reloaded.total_time == 42
        assert reloaded.scale_crop is True

    def it_leaves_unmapped_elements_untouched(self):
        """HeadingPairs and TitlesOfParts have no API but must survive a round trip."""
        document = docx.Document()
        document.extended_properties.company = "Acme Corp"

        stream = io.BytesIO()
        document.save(stream)

        import zipfile

        with zipfile.ZipFile(stream) as z:
            app_xml = z.read("docProps/app.xml").decode("utf-8")
        assert "HeadingPairs" in app_xml
        assert "TitlesOfParts" in app_xml
