# pyright: reportPrivateUsage=false

"""Unit test suite for the VBA-project accessors on |Document|.

The `vbaProject.bin` in `tests/test_files/macro-enabled.docm` is a 27-byte stub whose
header is not a valid OLE signature. That is sufficient here — nothing in this tier
parses the blob — but nothing may be built on top of it as though it were a real
project.
"""

from __future__ import annotations

import io

import docx
from docx.opc.constants import CONTENT_TYPE as CT

from .unitutil.file import test_file


def _macro_enabled():
    return docx.Document(test_file("macro-enabled.docm"))


def _reopened(document) -> docx.document.Document:
    stream = io.BytesIO()
    document.save(stream)
    return docx.Document(io.BytesIO(stream.getvalue()))


class DescribeVbaProjectAccess:
    """Unit-test suite for `Document.vba_project` and `Document.has_macros`."""

    def it_reads_the_project_bytes_of_a_macro_enabled_document(self):
        document = _macro_enabled()

        assert document.has_macros is True
        assert isinstance(document.vba_project, bytes)
        assert len(document.vba_project) == 27

    def and_reports_None_for_a_document_with_no_macros(self):
        document = docx.Document()

        assert document.has_macros is False
        assert document.vba_project is None

    def it_round_trips_the_project_untouched(self):
        document = _macro_enabled()
        before = document.vba_project

        assert _reopened(document).vba_project == before


class DescribeVbaProjectTransplant:
    """Assigning a project must switch the main part's content type with it."""

    def it_adds_the_project_and_switches_the_content_type(self):
        donor = _macro_enabled()
        document = docx.Document()

        document.vba_project = donor.vba_project

        assert document.has_macros is True
        assert document.part.content_type == CT.WML_DOCUMENT_MACRO_ENABLED_MAIN

    def and_the_transplanted_project_survives_a_save(self):
        donor = _macro_enabled()
        document = docx.Document()
        document.vba_project = donor.vba_project

        reopened = _reopened(document)

        assert reopened.vba_project == donor.vba_project
        assert reopened.part.content_type == CT.WML_DOCUMENT_MACRO_ENABLED_MAIN

    def it_replaces_an_existing_project_rather_than_adding_a_second(self):
        document = _macro_enabled()

        document.vba_project = b"replacement bytes"

        assert document.vba_project == b"replacement bytes"
        assert _reopened(document).vba_project == b"replacement bytes"

    def it_switches_a_template_to_the_macro_enabled_template_type(self):
        donor = _macro_enabled()
        document = docx.Document()
        document.save(io.BytesIO(), as_template=True)

        document.vba_project = donor.vba_project

        assert document.part.content_type == CT.WML_TEMPLATE_MACRO_ENABLED_MAIN


class DescribeVbaProjectRemoval:
    """Stripping the macros from a received document."""

    def it_removes_the_project_and_switches_the_content_type_back(self):
        document = _macro_enabled()

        removed = document.remove_vba_project()

        assert removed == 1
        assert document.has_macros is False
        assert document.part.content_type == CT.WML_DOCUMENT_MAIN

    def and_del_does_the_same(self):
        document = _macro_enabled()

        del document.vba_project

        assert document.has_macros is False
        assert document.part.content_type == CT.WML_DOCUMENT_MAIN

    def and_assigning_None_does_the_same(self):
        document = _macro_enabled()

        document.vba_project = None

        assert document.has_macros is False

    def it_produces_a_document_word_will_not_warn_about(self):
        """A main part claiming to be macro-enabled with no project makes Word warn."""
        document = _macro_enabled()
        del document.vba_project

        reopened = _reopened(document)

        assert reopened.has_macros is False
        assert reopened.part.content_type == CT.WML_DOCUMENT_MAIN
        assert "vbaProject.bin" not in [
            str(p.partname) for p in reopened.part.package.iter_parts()
        ]

    def it_reports_zero_for_a_document_with_no_project(self):
        document = docx.Document()

        assert document.remove_vba_project() == 0
        # -- and leaves the content type alone --
        assert document.part.content_type == CT.WML_DOCUMENT_MAIN
