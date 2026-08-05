# pyright: reportPrivateUsage=false

"""Unit test suite for defining numbering from scratch — the write side of #98."""

from __future__ import annotations

import io

import pytest

import docx
from docx.enum.numbering import WD_NUMBER_FORMAT
from docx.oxml.ns import qn
from docx.oxml.numbering import CT_AbstractNum, CT_Lvl
from docx.shared import Inches


def _reopened(document):
    stream = io.BytesIO()
    document.save(stream)
    return docx.Document(io.BytesIO(stream.getvalue()))


class DescribeCT_LvlWriteSide:
    """Unit-test suite for the setters added to `w:lvl`."""

    def it_writes_each_child_in_schema_order(self):
        """`CT_Lvl` is an `xsd:sequence`; Word rejects the children out of order."""
        lvl = CT_Lvl.new(0)

        # -- deliberately assigned back to front --
        lvl.jc = "left"
        lvl.lvl_text = "%1."
        lvl.suffix = "space"
        lvl.p_style = "ListNumber"
        lvl.lvl_restart = 0
        lvl.num_fmt = WD_NUMBER_FORMAT.DECIMAL
        lvl.start = 1

        tags = [child.tag.split("}")[1] for child in lvl]
        assert tags == ["start", "numFmt", "lvlRestart", "pStyle", "suff", "lvlText", "lvlJc"]

    def it_reads_back_what_it_wrote(self):
        lvl = CT_Lvl.new(3)

        lvl.start = 5
        lvl.num_fmt = WD_NUMBER_FORMAT.LOWER_ROMAN
        lvl.lvl_text = "%4)"
        lvl.suffix = "nothing"
        lvl.lvl_restart = 2
        lvl.p_style = "Quote"
        lvl.is_lgl = True

        assert lvl.ilvl == 3
        assert lvl.start == 5
        assert lvl.num_fmt == WD_NUMBER_FORMAT.LOWER_ROMAN
        assert lvl.lvl_text == "%4)"
        assert lvl.suffix == "nothing"
        assert lvl.lvl_restart == 2
        assert lvl.p_style == "Quote"
        assert lvl.is_lgl is True

    def it_removes_a_child_assigned_None(self):
        lvl = CT_Lvl.new(0)
        lvl.start = 1

        lvl.start = None

        assert lvl.start is None
        assert lvl.find(qn("w:start")) is None

    def it_rejects_an_invalid_suffix(self):
        lvl = CT_Lvl.new(0)

        with pytest.raises(ValueError, match="suffix must be one of"):
            lvl.suffix = "sideways"

    def it_accepts_a_number_format_as_a_string(self):
        """Word admits vendor extensions here, so an unlisted value must be writable."""
        lvl = CT_Lvl.new(0)

        lvl.num_fmt = "custom"

        assert lvl.find(qn("w:numFmt")).get(qn("w:val")) == "custom"


class DescribeCT_AbstractNumWriteSide:
    """Unit-test suite for building an abstract definition."""

    def it_keeps_levels_in_ascending_ilvl_order(self):
        """Word rejects an abstract definition whose levels are out of order."""
        abstract = CT_AbstractNum.new(0)

        for ilvl in (2, 0, 3, 1):
            abstract.add_level(ilvl)

        assert [lvl.ilvl for lvl in abstract.lvl_lst] == [0, 1, 2, 3]

    def it_returns_an_existing_level_rather_than_adding_a_second(self):
        abstract = CT_AbstractNum.new(0)

        first = abstract.add_level(0)
        again = abstract.add_level(0)

        assert first is again
        assert len(abstract.lvl_lst) == 1

    def it_puts_multiLevelType_before_the_levels(self):
        abstract = CT_AbstractNum.new(0)
        abstract.add_level(0)

        abstract.multi_level_type = "multilevel"

        tags = [child.tag.split("}")[1] for child in abstract]
        assert tags == ["multiLevelType", "lvl"]

    def it_rejects_an_invalid_multi_level_type(self):
        abstract = CT_AbstractNum.new(0)

        with pytest.raises(ValueError, match="multi_level_type must be one of"):
            abstract.multi_level_type = "sometimes"

    def it_writes_no_nsid_or_tmpl(self):
        """Those are what Word uses to claim a definition as one of its own."""
        abstract = CT_AbstractNum.new(0)
        abstract.multi_level_type = "multilevel"

        assert abstract.find(qn("w:nsid")) is None
        assert abstract.find(qn("w:tmpl")) is None


class DescribeAddDefinition:
    """Unit-test suite for `Numbering.add_definition()`."""

    def it_defines_a_list_the_template_does_not_have(self):
        document = docx.Document()

        definition = document.numbering.add_definition(
            [
                {"number_format": WD_NUMBER_FORMAT.DECIMAL, "level_text": "%1)"},
                {"number_format": WD_NUMBER_FORMAT.LOWER_LETTER, "level_text": "%2)"},
            ]
        )

        assert definition.level(0).level_text == "%1)"
        assert definition.level(1).number_format == WD_NUMBER_FORMAT.LOWER_LETTER

    def and_the_list_numbers_as_defined(self):
        document = docx.Document()
        definition = document.numbering.add_definition(
            [
                {"number_format": WD_NUMBER_FORMAT.DECIMAL, "level_text": "%1)"},
                {"number_format": WD_NUMBER_FORMAT.LOWER_LETTER, "level_text": "%2)"},
            ]
        )
        for text, level in (("one", 0), ("sub", 1), ("two", 0)):
            document.add_paragraph(text).set_numbering(definition.num_id, level=level)

        assert [p.list_number for p in document.paragraphs] == ["1)", "a)", "2)"]

    def and_it_survives_a_save(self):
        document = docx.Document()
        definition = document.numbering.add_definition(
            [{"number_format": WD_NUMBER_FORMAT.UPPER_ROMAN, "level_text": "%1."}]
        )
        document.add_paragraph("item").set_numbering(definition.num_id, level=0)

        reopened = _reopened(document)

        assert reopened.paragraphs[0].list_number == "I."

    def it_allocates_free_ids(self):
        document = docx.Document()
        existing_abstract = {
            a.abstractNumId
            for a in document.numbering._element.abstractNum_lst  # pyright: ignore
        }

        definition = document.numbering.add_definition([{}])

        assert definition.abstract_num_id not in existing_abstract
        assert document.numbering.get(definition.num_id) is not None

    def it_places_the_abstract_definition_before_every_num(self):
        """`CT_Numbering` is an `xsd:sequence`; Word rejects them the other way round."""
        document = docx.Document()

        document.numbering.add_definition([{}])

        tags = [child.tag.split("}")[1] for child in document.numbering._element]
        assert tags == sorted(tags, key=lambda t: 0 if t == "abstractNum" else 1)

    def it_defaults_to_nine_decimal_levels(self):
        document = docx.Document()

        definition = document.numbering.add_definition()

        assert definition.level(8).number_format == WD_NUMBER_FORMAT.DECIMAL
        assert definition.level(8).level_text == "%9."

    def it_writes_the_multi_level_type_when_given(self):
        document = docx.Document()

        definition = document.numbering.add_definition([{}], multi_level_type="singleLevel")

        abstract = document.numbering._element.abstractNum_having_abstractNumId(  # pyright: ignore
            definition.abstract_num_id
        )
        assert abstract.multi_level_type == "singleLevel"


class DescribeListShorthands:
    """Unit-test suite for the two common-case builders."""

    def it_builds_a_numbered_list_with_alternating_formats(self):
        document = docx.Document()

        definition = document.numbering.add_numbered_definition(3)

        assert [definition.level(i).number_format for i in range(3)] == [
            WD_NUMBER_FORMAT.DECIMAL,
            WD_NUMBER_FORMAT.LOWER_LETTER,
            WD_NUMBER_FORMAT.LOWER_ROMAN,
        ]

    def and_indents_each_level_further_than_the_one_above(self):
        document = docx.Document()

        definition = document.numbering.add_numbered_definition(3)

        indents = [definition.level(i).indent for i in range(3)]
        assert indents == [Inches(0.5), Inches(0.75), Inches(1.0)]
        assert definition.level(0).hanging_indent == Inches(0.25)

    def it_builds_a_bulleted_list(self):
        document = docx.Document()

        definition = document.numbering.add_bulleted_definition(3)

        assert all(definition.level(i).is_bullet for i in range(3))
        assert [definition.level(i).level_text for i in range(3)] == ["•", "o", "§"]

    def and_the_bullets_can_be_chosen(self):
        document = docx.Document()

        definition = document.numbering.add_bulleted_definition(2, bullets=("-", "*"))

        assert [definition.level(i).level_text for i in range(2)] == ["-", "*"]

    def a_bulleted_list_renders_its_bullet(self):
        document = docx.Document()
        definition = document.numbering.add_bulleted_definition(2)
        document.add_paragraph("item").set_numbering(definition.num_id, level=0)

        assert document.paragraphs[0].list_number == "•"


class DescribeNumberingLevelSet:
    """Unit-test suite for `NumberingLevel.set()`."""

    def it_changes_only_what_it_is_given(self):
        document = docx.Document()
        definition = document.numbering.add_definition(
            [{"number_format": WD_NUMBER_FORMAT.DECIMAL, "level_text": "%1.", "start": 3}]
        )

        definition.level(0).set(level_text="%1)")

        assert definition.level(0).level_text == "%1)"
        assert definition.level(0).start == 3
        assert definition.level(0).number_format == WD_NUMBER_FORMAT.DECIMAL

    def it_can_set_the_suffix(self):
        document = docx.Document()
        definition = document.numbering.add_definition([{}])

        definition.level(0).set(suffix="space")

        assert definition.level(0).suffix == "space"

    def and_the_suffix_defaults_to_tab(self):
        document = docx.Document()
        definition = document.numbering.add_definition([{}])

        assert definition.level(0).suffix == "tab"

    def it_can_set_the_indents(self):
        document = docx.Document()
        definition = document.numbering.add_definition([{}])

        definition.level(0).set(indent=Inches(1), hanging_indent=Inches(0.25))

        assert definition.level(0).indent == Inches(1)
        assert definition.level(0).hanging_indent == Inches(0.25)

    def it_returns_self_for_chaining(self):
        document = docx.Document()
        definition = document.numbering.add_definition([{}])
        level = definition.level(0)

        assert level.set(start=2) is level

    def it_raises_for_a_level_with_no_definition(self):
        """One of the nine levels the abstract definition does not define."""
        document = docx.Document()
        definition = document.numbering.add_definition([{}])

        with pytest.raises(ValueError, match="has no definition"):
            definition.level(5).set(start=1)
