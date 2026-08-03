"""Unit test suite for cross-run search and replace.

Covers `docx.oxml.text.isolate`, `docx.text.search`, and the `replace_text()` and
`isolate_run()` methods layered over them.
"""

from __future__ import annotations

import io
from typing import cast

import pytest

import docx
from docx import types as t
from docx.oxml.text.isolate import isolate_range
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph

from ..unitutil.cxml import element, xml


class DescribeIsolateRange:
    """Unit-test suite for `docx.oxml.text.isolate.isolate_range`."""

    @pytest.mark.parametrize(
        ("p_cxml", "start", "end", "expected_texts"),
        [
            # -- range is already exactly one run --
            ('w:p/(w:r/w:t"abc",w:r/w:t"def")', 0, 3, ["abc"]),
            ('w:p/(w:r/w:t"abc",w:r/w:t"def")', 3, 6, ["def"]),
            # -- range inside one run: the run is divided in three --
            ('w:p/w:r/w:t"abcdef"', 2, 4, ["cd"]),
            # -- range spanning two runs: one piece from each --
            ('w:p/(w:r/w:t"abc",w:r/w:t"def")', 2, 4, ["c", "d"]),
            # -- range covering everything --
            ('w:p/(w:r/w:t"ab",w:r/w:t"cd")', 0, 4, ["ab", "cd"]),
            # -- empty and out-of-range --
            ('w:p/w:r/w:t"abc"', 1, 1, []),
            ("w:p", 0, 3, []),
        ],
    )
    def it_splits_the_runs_covering_the_range(
        self, p_cxml: str, start: int, end: int, expected_texts: list[str]
    ):
        p = cast(CT_P, element(p_cxml))
        original_text = p.text

        runs = isolate_range(p, start, end)

        assert [r.text for r in runs] == expected_texts
        # -- splitting divides the text differently but never changes it --
        assert p.text == original_text

    def it_carries_the_run_formatting_onto_each_piece(self):
        p = cast(CT_P, element('w:p/w:r/(w:rPr/w:b,w:t"abcdef")'))

        isolate_range(p, 2, 4)

        assert p.xml == xml(
            'w:p/(w:r/(w:rPr/w:b,w:t"ab"),w:r/(w:rPr/w:b,w:t"cd"),w:r/(w:rPr/w:b,w:t"ef"))'
        )

    def it_preserves_significant_whitespace_when_dividing_a_run(self):
        """Without `xml:space` a parser may drop the space, joining two words."""
        p = cast(CT_P, element('w:p/w:r/w:t"alpha beta"'))

        isolate_range(p, 6, 10)

        assert p.text == "alpha beta"
        t_elms = p.xpath(".//w:t")
        assert t_elms[0].get("{http://www.w3.org/XML/1998/namespace}space") == "preserve", (
            "trailing space would be collapsed away"
        )

    def it_divides_runs_inside_a_hyperlink_without_moving_them_out(self):
        p = cast(CT_P, element('w:p/w:hyperlink/w:r/w:t"abcdef"'))

        runs = isolate_range(p, 2, 4)

        assert [r.text for r in runs] == ["cd"]
        assert len(p.xpath("./w:hyperlink/w:r")) == 3, "pieces stayed in the hyperlink"

    def it_treats_a_tab_as_one_indivisible_character(self):
        p = cast(CT_P, element('w:p/w:r/(w:t"a",w:tab,w:t"b")'))

        runs = isolate_range(p, 1, 2)

        assert [r.text for r in runs] == ["\t"]

    def it_ignores_a_field_instruction(self):
        """`w:instrText` is a field instruction, not document text."""
        p = cast(CT_P, element('w:p/(w:r/w:instrText" PAGE ",w:r/w:t"abc")'))

        runs = isolate_range(p, 0, 3)

        assert [r.text for r in runs] == ["abc"]

    @pytest.mark.parametrize(("start", "end"), [(-1, 3), (3, 1)])
    def it_raises_on_an_invalid_range(self, start: int, end: int):
        p = cast(CT_P, element('w:p/w:r/w:t"abcdef"'))

        with pytest.raises(ValueError, match="invalid character range"):
            isolate_range(p, start, end)


class DescribeParagraphIsolateRun:
    """Unit-test suite for `docx.text.paragraph.Paragraph.isolate_run`."""

    def it_returns_the_range_as_a_single_run(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(cast(CT_P, element('w:p/w:r/w:t"the important part"')), fake_parent)

        run = paragraph.isolate_run(4, 13)

        assert run.text == "important"
        assert paragraph.text == "the important part"

    def it_merges_runs_when_the_range_spans_several(self, fake_parent: t.ProvidesStoryPart):
        p = cast(CT_P, element('w:p/(w:r/w:t"the impor",w:r/w:t"tant part")'))
        paragraph = Paragraph(p, fake_parent)

        run = paragraph.isolate_run(4, 13)
        run.bold = True

        assert run.text == "important"
        assert paragraph.text == "the important part"
        assert [r.text for r in paragraph.runs] == ["the ", "important", " part"]
        assert [r.bold for r in paragraph.runs] == [None, True, None]

    def it_takes_the_formatting_of_the_run_at_the_start(self, fake_parent: t.ProvidesStoryPart):
        p = cast(CT_P, element('w:p/(w:r/(w:rPr/w:b,w:t"ab"),w:r/(w:rPr/w:i,w:t"cd"))'))
        paragraph = Paragraph(p, fake_parent)

        run = paragraph.isolate_run(0, 4)

        assert run.bold is True
        assert run.italic is None

    def it_raises_when_the_range_crosses_a_hyperlink_boundary(
        self, fake_parent: t.ProvidesStoryPart
    ):
        p = cast(CT_P, element('w:p/(w:r/w:t"ab",w:hyperlink/w:r/w:t"cd")'))
        paragraph = Paragraph(p, fake_parent)

        with pytest.raises(ValueError, match="spans a hyperlink"):
            paragraph.isolate_run(1, 3)

    def it_raises_on_an_empty_range(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(cast(CT_P, element('w:p/w:r/w:t"abc"')), fake_parent)

        with pytest.raises(ValueError, match="empty or lies beyond"):
            paragraph.isolate_run(1, 1)


class DescribeParagraphReplaceText:
    """Unit-test suite for `docx.text.paragraph.Paragraph.replace_text`."""

    @pytest.mark.parametrize(
        ("p_cxml", "old", "new", "expected_count", "expected_text"),
        [
            # -- the whole point: a match Word split across runs --
            ('w:p/(w:r/w:t"Hello wo",w:r/w:t"rld!")', "world", "there", 1, "Hello there!"),
            # -- three runs, match spanning all three --
            ('w:p/(w:r/w:t"a",w:r/w:t"bc",w:r/w:t"d")', "abcd", "X", 1, "X"),
            # -- several matches --
            ('w:p/w:r/w:t"a a a"', "a", "b", 3, "b b b"),
            # -- no match --
            ('w:p/w:r/w:t"abc"', "z", "y", 0, "abc"),
            # -- deletion --
            ('w:p/w:r/w:t"abcdef"', "cd", "", 1, "abef"),
            # -- empty paragraph --
            ("w:p", "a", "b", 0, ""),
        ],
    )
    def it_replaces_text_across_runs(
        self,
        p_cxml: str,
        old: str,
        new: str,
        expected_count: int,
        expected_text: str,
        fake_parent: t.ProvidesStoryPart,
    ):
        paragraph = Paragraph(cast(CT_P, element(p_cxml)), fake_parent)

        count = paragraph.replace_text(old, new)

        assert count == expected_count
        assert paragraph.text == expected_text

    def it_gives_the_replacement_the_formatting_of_the_match_start(
        self, fake_parent: t.ProvidesStoryPart
    ):
        p = cast(CT_P, element('w:p/(w:r/(w:rPr/w:b,w:t"keep bo"),w:r/w:t"ld here")'))
        paragraph = Paragraph(p, fake_parent)

        paragraph.replace_text("bold", "BOLD")

        assert paragraph.text == "keep BOLD here"
        bold_run = next(r for r in paragraph.runs if "BOLD" in r.text)
        assert bold_run.bold is True

    def it_leaves_the_surrounding_run_formatting_alone(self, fake_parent: t.ProvidesStoryPart):
        p = cast(CT_P, element('w:p/(w:r/(w:rPr/w:b,w:t"A"),w:r/w:t"x",w:r/(w:rPr/w:i,w:t"B"))'))
        paragraph = Paragraph(p, fake_parent)

        paragraph.replace_text("x", "y")

        assert paragraph.text == "AyB"
        assert [(r.text, r.bold, r.italic) for r in paragraph.runs] == [
            ("A", True, None),
            ("y", None, None),
            ("B", None, True),
        ]

    def it_keeps_a_hyperlink_intact_when_the_match_crosses_its_boundary(
        self, fake_parent: t.ProvidesStoryPart
    ):
        p = cast(CT_P, element('w:p/(w:r/w:t"see py",w:hyperlink/w:r/w:t"thon docs")'))
        paragraph = Paragraph(p, fake_parent)

        count = paragraph.replace_text("python", "PY")

        assert count == 1
        assert paragraph.text == "see PY docs"
        assert len(p.xpath("./w:hyperlink")) == 1, "the hyperlink itself survived"
        assert p.xpath("./w:hyperlink")[0].text == " docs"

    def it_replaces_text_inside_a_content_control(self, fake_parent: t.ProvidesStoryPart):
        p = cast(CT_P, element('w:p/w:sdt/w:sdtContent/w:r/w:t"Enter name"'))
        paragraph = Paragraph(p, fake_parent)

        count = paragraph.replace_text("Enter name", "Ada")

        assert count == 1
        assert paragraph.text == "Ada"

    def it_clears_the_placeholder_flag_of_a_control_it_writes_into(
        self, fake_parent: t.ProvidesStoryPart
    ):
        """Word discards text written into a control still marked as showing a prompt."""
        p = cast(
            CT_P,
            element(
                "w:p/w:sdt/(w:sdtPr/w:showingPlcHdr,"
                'w:sdtContent/w:r/w:t"Click here to enter text.")'
            ),
        )
        paragraph = Paragraph(p, fake_parent)

        paragraph.replace_text("Click here to enter text.", "Ada")

        assert paragraph.text == "Ada"
        assert p.xpath(".//w:showingPlcHdr") == []

    def it_never_matches_a_field_instruction(self, fake_parent: t.ProvidesStoryPart):
        """Editing `w:instrText` breaks the field; it is not document text."""
        p = cast(CT_P, element('w:p/(w:r/w:instrText" PAGE ",w:r/w:t"PAGE here")'))
        paragraph = Paragraph(p, fake_parent)

        count = paragraph.replace_text("PAGE", "X")

        assert count == 1
        assert p.xpath(".//w:instrText")[0].text == " PAGE "

    def it_matches_the_search_text_literally_by_default(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(cast(CT_P, element('w:p/w:r/w:t"axb a.b"')), fake_parent)

        count = paragraph.replace_text("a.b", "X")

        assert count == 1
        assert paragraph.text == "axb X"

    def it_can_match_a_regular_expression_with_group_references(
        self, fake_parent: t.ProvidesStoryPart
    ):
        p = cast(CT_P, element('w:p/w:r/w:t"2024-01-02"'))
        paragraph = Paragraph(p, fake_parent)

        count = paragraph.replace_text(r"(\d{4})-(\d{2})-(\d{2})", r"\3/\2/\1", regex=True)

        assert count == 1
        assert paragraph.text == "02/01/2024"

    def it_honors_the_count_limit(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(cast(CT_P, element('w:p/w:r/w:t"a a a a"')), fake_parent)

        count = paragraph.replace_text("a", "b", count=2)

        assert count == 2
        assert paragraph.text == "b b a a"

    def it_terminates_when_the_replacement_re_matches_the_pattern(
        self, fake_parent: t.ProvidesStoryPart
    ):
        """`replace("a", "aa")` must not feed on its own output."""
        paragraph = Paragraph(cast(CT_P, element('w:p/w:r/w:t"aaa"')), fake_parent)

        count = paragraph.replace_text("a", "aa")

        assert count == 3
        assert paragraph.text == "aaaaaa"

    def it_terminates_on_a_zero_width_match(self, fake_parent: t.ProvidesStoryPart):
        paragraph = Paragraph(cast(CT_P, element('w:p/w:r/w:t"ab"')), fake_parent)

        count = paragraph.replace_text("", "-", regex=True)

        assert count == 3
        assert paragraph.text == "-a-b-"

    def it_translates_tabs_and_newlines_in_the_replacement(self, fake_parent: t.ProvidesStoryPart):
        p = cast(CT_P, element('w:p/w:r/w:t"a|b"'))
        paragraph = Paragraph(p, fake_parent)

        paragraph.replace_text("|", "\t")

        assert paragraph.text == "a\tb"
        assert len(p.xpath(".//w:tab")) == 1, "a tab is an element, not a character"


class DescribeContainerReplaceText:
    """Unit-test suite for `replace_text()` on containers and the document."""

    def it_replaces_text_in_tables_including_nested_ones(self):
        document = docx.Document()
        table = document.add_table(rows=1, cols=1)
        inner = table.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "X marks it"

        count = document.replace_text("X", "Y")

        assert count == 1
        assert inner.cell(0, 0).text == "Y marks it"

    def it_can_be_told_to_skip_tables(self):
        document = docx.Document()
        document.add_paragraph("X")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "X"

        count = document.replace_text("X", "Y", tables=False)

        assert count == 1
        assert document.paragraphs[0].text == "Y"
        assert document.tables[0].cell(0, 0).text == "X"

    def it_leaves_headers_and_footnotes_alone_by_default(self):
        document = docx.Document()
        document.add_paragraph("{{x}}")
        document.sections[0].header.paragraphs[0].text = "{{x}}"
        footnote = document.footnotes.add_footnote("{{x}}")

        count = document.replace_text("{{x}}", "y")

        assert count == 1
        assert document.sections[0].header.paragraphs[0].text == "{{x}}"
        assert footnote.text == "{{x}}"

    def it_searches_headers_footers_and_footnotes_when_asked(self):
        document = docx.Document()
        document.add_paragraph("{{x}}")
        document.sections[0].header.paragraphs[0].text = "{{x}}"
        document.sections[0].footer.paragraphs[0].text = "{{x}}"
        footnote = document.footnotes.add_footnote("{{x}}")

        count = document.replace_text("{{x}}", "y", headers_footers=True, footnotes=True)

        assert count == 4
        assert document.sections[0].header.paragraphs[0].text == "y"
        assert document.sections[0].footer.paragraphs[0].text == "y"
        assert footnote.text == "y"

    def it_does_not_create_a_footnotes_part_just_to_search_it(self):
        document = docx.Document()
        document.add_paragraph("x")

        document.replace_text("x", "y", footnotes=True)

        assert document.part.has_footnotes_part is False

    def it_honors_the_count_limit_across_the_whole_document(self):
        document = docx.Document()
        for _ in range(4):
            document.add_paragraph("x")

        count = document.replace_text("x", "y", count=3)

        assert count == 3
        assert [p.text for p in document.paragraphs] == ["y", "y", "y", "x"]

    def it_survives_a_round_trip(self):
        document = docx.Document()
        document.add_paragraph("Dear {{name}},")

        document.replace_text("{{name}}", "Ada")
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        assert docx.Document(stream).paragraphs[0].text == "Dear Ada,"
