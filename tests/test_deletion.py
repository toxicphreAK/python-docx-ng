# pyright: reportPrivateUsage=false

"""Unit test suite for the deletion API.

Deletion is half of editing, and the reason to do it properly rather than bless the
`p._element.getparent().remove(p._element)` one-liner is everything the removed content
referred to.
"""

from __future__ import annotations

import io

import pytest

import docx


class DescribeParagraphDeletion:
    """Unit-test suite for `Paragraph.delete()`."""

    def it_can_delete_a_paragraph(self):
        document = docx.Document()
        for text in ("one", "two", "three"):
            document.add_paragraph(text)

        document.paragraphs[1].delete()

        assert [p.text for p in document.paragraphs] == ["one", "three"]

    def it_drops_a_relationship_the_deleted_content_alone_referenced(self):
        """A relationship to nothing is what makes a document need repairing."""
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_hyperlink("link", "https://example.com/")
        assert _external_rel_count(document) == 1

        paragraph.delete()

        assert _external_rel_count(document) == 0

    def it_keeps_a_relationship_still_referenced_elsewhere(self):
        document = docx.Document()
        first = document.add_paragraph()
        first.add_hyperlink("one", "https://example.com/")
        second = document.add_paragraph()
        second.add_hyperlink("two", "https://example.com/")

        first.delete()

        assert _external_rel_count(document) == 1
        assert second.hyperlinks[0].address == "https://example.com/"

    def it_removes_the_other_half_of_a_bookmark_it_breaks(self):
        """An unmatched delimiter left behind points at content that is gone."""
        document = docx.Document()
        first = document.add_paragraph("first").runs[0]
        last = document.add_paragraph("last").runs[0]
        first.mark_bookmark_range(last, "Span")
        assert len(document.bookmarks) == 1

        document.paragraphs[1].delete()

        assert len(document.bookmarks) == 0
        assert document.element.xpath("//w:bookmarkStart") == []

    def it_refuses_to_delete_the_only_paragraph_of_a_table_cell(self):
        """A `w:tc` without a block-level child is a document Word will not open."""
        document = docx.Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "content"

        with pytest.raises(ValueError, match="only block-level element"):
            cell.paragraphs[0].delete()

    def it_can_delete_a_paragraph_from_a_cell_that_has_another(self):
        document = docx.Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "first"
        cell.add_paragraph("second")

        cell.paragraphs[0].delete()

        assert [p.text for p in cell.paragraphs] == ["second"]

    def it_leaves_a_loadable_document(self):
        document = docx.Document()
        document.add_paragraph("keep")
        document.add_paragraph("drop").delete()

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        assert [p.text for p in docx.Document(stream).paragraphs] == ["keep"]


class DescribeRunDeletion:
    """Unit-test suite for `Run.delete()`."""

    def it_can_delete_a_run(self):
        paragraph = docx.Document().add_paragraph()
        paragraph.add_run("a")
        paragraph.add_run("b")

        paragraph.runs[0].delete()

        assert paragraph.text == "b"


class DescribeTableDeletion:
    """Unit-test suite for deleting rows, columns and whole tables."""

    def it_can_delete_a_row(self, table_3x3: docx.table.Table):
        table_3x3.rows[1].delete()

        assert _grid(table_3x3) == [["00", "01", "02"], ["20", "21", "22"]]

    def it_can_delete_a_column(self, table_3x3: docx.table.Table):
        table_3x3.columns[1].delete()

        assert len(table_3x3.columns) == 2
        assert _grid(table_3x3) == [["00", "02"], ["10", "12"], ["20", "22"]]

    def it_narrows_a_spanning_cell_rather_than_removing_it(self):
        """The rest of the span has to survive removing one column of it."""
        document = docx.Document()
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "wide"

        table.columns[0].delete()

        assert len(table.columns) == 2
        assert table.cell(0, 0).text == "wide"
        assert table.cell(0, 0).grid_span == 1

    def it_hands_a_vertical_span_to_the_row_below_when_its_origin_row_goes(self):
        document = docx.Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).merge(table.cell(2, 0))
        table.cell(0, 0).text = "M"

        table.rows[0].delete()

        assert table.cell(0, 0).text == "M"
        assert table.cell(0, 0).span == (2, 1)

    def it_can_delete_a_whole_table(self):
        document = docx.Document()
        table = document.add_table(rows=2, cols=2)
        document.add_table(rows=1, cols=1)

        table.delete()

        assert len(document.tables) == 1

    def it_leaves_a_loadable_document(self, table_3x3: docx.table.Table):
        document = table_3x3.part.document
        table_3x3.columns[1].delete()
        table_3x3.rows[0].delete()

        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)

        reloaded = docx.Document(stream)
        assert _grid(reloaded.tables[0]) == [["10", "12"], ["20", "22"]]

    # fixtures ---------------------------------------------

    @pytest.fixture
    def table_3x3(self) -> docx.table.Table:
        document = docx.Document()
        table = document.add_table(rows=3, cols=3)
        table.style = None
        for row_idx in range(3):
            for col_idx in range(3):
                table.cell(row_idx, col_idx).text = "%d%d" % (row_idx, col_idx)
        return table


def _external_rel_count(document: docx.Document) -> int:
    return sum(1 for rel in document.part.rels.values() if rel.is_external)


def _grid(table: docx.table.Table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]
