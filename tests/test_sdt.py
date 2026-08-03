# pyright: reportPrivateUsage=false

"""Unit test suite for the docx.sdt module."""

from __future__ import annotations

from typing import cast

import pytest

from docx.blkcntnr import BlockItemContainer
from docx.enum.text import WD_CONTENT_CONTROL_TYPE
from docx.oxml.document import CT_Body
from docx.oxml.sdt import CT_Sdt
from docx.oxml.text.paragraph import CT_P
from docx.sdt import ContentControl
from docx.text.paragraph import Paragraph

from .unitutil.cxml import element


class DescribeContentControl:
    """Unit-test suite for `docx.sdt.ContentControl`."""

    def it_knows_its_tag_alias_id_and_type(self):
        sdt = cast(
            CT_Sdt,
            element(
                "w:sdt/w:sdtPr/(w:alias{w:val=Client name},w:tag{w:val=client},"
                "w:id{w:val=42},w:text)"
            ),
        )

        content_control = ContentControl(sdt, None)

        assert content_control.alias == "Client name"
        assert content_control.tag == "client"
        assert content_control.id == 42
        assert content_control.type == WD_CONTENT_CONTROL_TYPE.TEXT

    def it_provides_access_to_the_paragraphs_it_wraps(self):
        sdt = cast(
            CT_Sdt, element('w:sdt/w:sdtContent/(w:p/w:r/w:t"one",w:p/w:r/w:t"two")')
        )

        content_control = ContentControl(sdt, None)

        assert [p.text for p in content_control.paragraphs] == ["one", "two"]
        assert content_control.is_block_level is True

    def it_provides_access_to_the_table_it_wraps(self):
        sdt = cast(
            CT_Sdt,
            element(
                "w:sdt/w:sdtContent/w:tbl/(w:tblPr,w:tblGrid/w:gridCol,"
                'w:tr/w:tc/w:p/w:r/w:t"cell")'
            ),
        )

        content_control = ContentControl(sdt, None)

        assert len(content_control.tables) == 1
        assert content_control.tables[0].cell(0, 0).text == "cell"
        assert [type(i).__name__ for i in content_control.iter_inner_content()] == ["Table"]

    def it_provides_access_to_the_runs_of_a_run_level_control(self):
        sdt = cast(CT_Sdt, element('w:sdt/w:sdtContent/(w:r/w:t"a",w:r/w:t"b")'))

        content_control = ContentControl(sdt, None)

        assert [r.text for r in content_control.runs] == ["a", "b"]
        assert content_control.paragraphs == []
        assert content_control.is_block_level is False

    def it_knows_the_text_it_contains(self):
        sdt = cast(
            CT_Sdt, element('w:sdt/w:sdtContent/(w:p/w:r/w:t"one",w:p/w:r/w:t"two")')
        )

        assert ContentControl(sdt, None).text == "one\ntwo"

    def it_knows_when_it_is_showing_its_placeholder(self):
        sdt = cast(CT_Sdt, element("w:sdt/w:sdtPr/w:showingPlcHdr"))

        assert ContentControl(sdt, None).showing_placeholder is True


class DescribeContentControlsInAContainer:
    """A container flattens control content but keeps the control itself reachable."""

    def it_includes_wrapped_paragraphs_in_the_container_paragraphs(self, body_cxml: str):
        container = BlockItemContainer(cast(CT_Body, element(body_cxml)), None)

        assert [p.text for p in container.paragraphs] == [
            "INSIDE",
            "OUTSIDE",
            "NESTED",
        ]

    def it_includes_wrapped_tables_in_the_container_tables(self, body_cxml: str):
        container = BlockItemContainer(cast(CT_Body, element(body_cxml)), None)

        assert len(container.tables) == 1

    def it_orders_flattened_content_as_it_appears_in_the_document(self, body_cxml: str):
        container = BlockItemContainer(cast(CT_Body, element(body_cxml)), None)

        assert [type(i).__name__ for i in container.iter_inner_content()] == [
            "Paragraph",
            "Paragraph",
            "Paragraph",
            "Table",
        ]

    def it_keeps_the_control_boundary_discoverable(self, body_cxml: str):
        """Flattening is what most callers want, but not all of them."""
        container = BlockItemContainer(cast(CT_Body, element(body_cxml)), None)

        # -- outermost first, so a nested control follows the one containing it --
        assert [cc.tag for cc in container.content_controls] == ["one", "outer", "inner"]

    # fixtures ---------------------------------------------

    @pytest.fixture
    def body_cxml(self) -> str:
        return (
            "w:body/("
            'w:sdt/(w:sdtPr/w:tag{w:val=one},w:sdtContent/w:p/w:r/w:t"INSIDE"),'
            'w:p/w:r/w:t"OUTSIDE",'
            "w:sdt/(w:sdtPr/w:tag{w:val=outer},w:sdtContent/("
            'w:sdt/(w:sdtPr/w:tag{w:val=inner},w:sdtContent/w:p/w:r/w:t"NESTED"),'
            "w:tbl/(w:tblPr,w:tblGrid/w:gridCol,w:tr/w:tc/w:p)))"
            ")"
        )


class DescribeContentControlsInAParagraph:
    """A run-level control is flattened into `.runs` but stays reachable."""

    def it_includes_wrapped_runs_in_the_paragraph_runs(self):
        p = cast(
            CT_P,
            element(
                'w:p/(w:r/w:t"a",w:sdt/(w:sdtPr/w:tag{w:val=inline},'
                'w:sdtContent/w:r/w:t"b"),w:r/w:t"c")'
            ),
        )

        paragraph = Paragraph(p, None)

        assert [r.text for r in paragraph.runs] == ["a", "b", "c"]
        assert paragraph.text == "abc"

    def it_keeps_the_control_boundary_discoverable(self):
        p = cast(
            CT_P,
            element(
                'w:p/(w:r/w:t"a",w:sdt/(w:sdtPr/w:tag{w:val=inline},'
                'w:sdtContent/w:r/w:t"b"))'
            ),
        )

        paragraph = Paragraph(p, None)

        assert [cc.tag for cc in paragraph.content_controls] == ["inline"]
        assert paragraph.content_controls[0].text == "b"
